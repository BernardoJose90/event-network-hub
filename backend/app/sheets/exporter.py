"""Export the live sheet model to .xlsx.

Calculated cells become real formulas over their current ranges, so the file
keeps working in Excel, Google Sheets and LibreOffice (which also makes our PDF).
"""

from __future__ import annotations

import io
from collections import defaultdict

from openpyxl import Workbook as XlsxWorkbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation

from .engine import compute, section_rows
from .model import EXPR_REF, Sheet, Source, Workbook

# Excel limits a literal dropdown list to 255 characters.
MAX_LITERAL_LIST = 255
# A total with nothing to add up yet. N("…") evaluates to 0 in Excel, Google
# Sheets and LibreOffice, and tells the importer what the cell is.
EMPTY_SECTION_TOTAL = '=N("section total")'
EMPTY_GRAND_TOTAL = '=N("grand total")'


def _width_chars(px: int) -> float:
    return max(1.0, round((px - 5) / 7, 2))


def _validation_formula(source: Source, workbook: Workbook) -> str | None:
    if source.values is not None:
        literal = ",".join(v.replace('"', "'") for v in source.values)
        return f'"{literal}"' if len(literal) + 2 <= MAX_LITERAL_LIST else None
    target = next((s for s in workbook.sheets if s.id == source.sheet), None)
    if target is None:
        return None
    col_index = next((i for i, c in enumerate(target.columns, 1) if c.id == source.column), None)
    if col_index is None:
        return None
    letter = get_column_letter(col_index)
    return f"{quote_sheetname(target.name)}!${letter}$1:${letter}${max(len(target.rows), 1)}"


def _write_sheet(ws, sheet: Sheet, workbook: Workbook) -> None:
    col_index = {c.id: i for i, c in enumerate(sheet.columns, 1)}
    row_index = {r.id: i for i, r in enumerate(sheet.rows, 1)}
    computed = compute(sheet)
    validations: dict[str, list[str]] = defaultdict(list)

    for c in sheet.columns:
        ws.column_dimensions[get_column_letter(col_index[c.id])].width = _width_chars(c.width)

    for r_idx, row in enumerate(sheet.rows, 1):
        for col_id, cell in row.cells.items():
            c_idx = col_index[col_id]
            letter = get_column_letter(c_idx)
            target = ws.cell(r_idx, c_idx)
            if cell.kind == "section_total":
                rows = section_rows(sheet, r_idx - 1, col_id)
                target.value = f"=SUM({letter}{rows.start + 1}:{letter}{rows.stop})" if len(rows) else EMPTY_SECTION_TOTAL
            elif cell.kind == "grand_total":
                refs = [
                    f"{letter}{i}" for i, r in enumerate(sheet.rows, 1)
                    if (c := r.cells.get(col_id)) is not None and c.kind == "section_total"
                ]
                target.value = f"=SUM({','.join(refs)})" if refs else EMPTY_GRAND_TOTAL
            elif cell.kind == "row_formula":
                missing = [ref for ref in EXPR_REF.findall(cell.expr or "") if ref not in col_index]
                if missing:
                    target.value = computed.get((row.id, col_id))
                else:
                    target.value = "=" + EXPR_REF.sub(
                        lambda m: f"{get_column_letter(col_index[m.group(1)])}{r_idx}", cell.expr or ""
                    )
            elif cell.kind == "checkbox":
                target.value = bool(cell.v)
            else:
                target.value = cell.v

            if cell.kind == "dropdown" and cell.source is not None:
                formula = _validation_formula(cell.source, workbook)
                if formula:
                    validations[formula].append(target.coordinate)

            s = cell.style
            target.font = Font(
                name=(s.font if s and s.font else sheet.font),
                size=(s.size if s and s.size else sheet.font_size),
                bold=bool(s and s.bold), italic=bool(s and s.italic),
                color=("FF" + s.color[1:]) if s and s.color else None,
            )
            if s is not None:
                if s.fill:
                    target.fill = PatternFill(fill_type="solid", fgColor="FF" + s.fill[1:])
                if s.align or s.wrap:
                    target.alignment = Alignment(horizontal=s.align, wrap_text=bool(s.wrap), vertical="top" if s.wrap else None)

    for formula, coords in validations.items():
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        for coord in coords:
            dv.add(coord)
        ws.add_data_validation(dv)

    for m in sheet.merges:
        if m.row in row_index and m.col in col_index:
            start = col_index[m.col]
            end = min(start + m.span - 1, len(sheet.columns))
            if end > start:
                ws.merge_cells(start_row=row_index[m.row], start_column=start,
                               end_row=row_index[m.row], end_column=end)

    if sheet.frozen_rows:
        ws.freeze_panes = f"A{sheet.frozen_rows + 1}"


def export_workbook(workbook: Workbook) -> bytes:
    xlsx = XlsxWorkbook()
    xlsx.remove(xlsx.active)
    for sheet in workbook.sheets:
        # Excel forbids some characters and > 31 chars in sheet titles.
        title = "".join("_" if ch in '[]:*?/\\' else ch for ch in sheet.name)[:31]
        _write_sheet(xlsx.create_sheet(title=title), sheet, workbook)
    out = io.BytesIO()
    xlsx.save(out)
    return out.getvalue()
