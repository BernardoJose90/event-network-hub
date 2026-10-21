from __future__ import annotations

import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import Settings
from app.main import create_app
from app.storage import version_key

from .conftest import TEST_DATABASE_URL, FakeConverter
from .records_fixture import records_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ADMIN = {"X-Role": "network-admin", "X-User": "alice"}
VIEWER = {"X-Role": "viewer", "X-User": "victor"}


@pytest.fixture
def client(tmp_path: Path):
    settings = Settings(data_dir=tmp_path, gotenberg_url="http://unused", gotenberg_timeout_seconds=1,
                        max_upload_bytes=1024 * 1024, database_url=TEST_DATABASE_URL)
    app = create_app(settings)
    with TestClient(app) as c:
        app.state.converter = FakeConverter()
        yield c


@pytest.fixture
def event(client: TestClient) -> int:
    r = client.post("/api/documents/upload", headers=ADMIN,
                    files={"file": ("LIS26 _ Event_Network_Requirements_V1.xlsx", records_workbook(), XLSX)})
    doc = r.json()["document"]["id"]
    assert client.post(f"/api/events/{doc}/import", headers=ADMIN, json={}).status_code == 200
    return doc


def _records(client, doc, rtype):
    return client.get(f"/api/events/{doc}/records/{rtype}", headers=ADMIN).json()


def _find(client, doc, rtype, **match):
    return next(r for r in _records(client, doc, rtype) if all(r["data"].get(k) == v for k, v in match.items()))


def _patch(client, doc, record, data):
    r = client.patch(f"/api/events/{doc}/records/{record['type']}/{record['id']}", headers=ADMIN,
                     json={"revision": record["revision"], "data": data})
    assert r.status_code == 200, r.text
    return r.json()


def _log(client, doc, **params):
    r = client.get(f"/api/events/{doc}/changelog", headers=VIEWER, params=params)
    assert r.status_code == 200, r.text
    return r.json()


def _section(log, rtype):
    return next(s for s in log["sections"] if s["type"] == rtype)


def test_first_version_and_draft(client: TestClient, event: int) -> None:
    draft = _log(client, event)
    assert draft["first"] and draft["from_version"] is None and draft["to_version"] == "draft"
    assert draft["items"] == 14 and draft["published_versions"] == []
    client.post(f"/api/events/{event}/publish", headers=ADMIN)
    assert _log(client, event)["counts"] == {"added": 0, "removed": 0, "changed": 0}
    v2 = _log(client, event, to=2)  # compared with V1, the uploaded workbook the records came from
    assert not v2["first"] and v2["from_version"] == 1 and v2["to_version"] == 2
    assert v2["counts"] == {"added": 0, "removed": 0, "changed": 0}
    assert v2["notices"] == ["V1 was uploaded, so its items were read from its Excel file."]


def test_net_changes_between_versions(client: TestClient, event: int) -> None:
    client.post(f"/api/events/{event}/publish", headers=ADMIN)  # V2

    office = _find(client, event, "ethernet", purpose="Photographers Office")
    office = _patch(client, event, office, {"ports": 4, "notes": "Two desks"})
    _patch(client, event, office, {"notes": None})  # undone: only the ports change remains
    prod = _find(client, event, "wifi", ssid="#Production")
    _patch(client, event, prod, {"password": "NewPass2!"})
    bar = _find(client, event, "ethernet", purpose="FoodTrucks / Bar")
    client.delete(f"/api/events/{event}/records/ethernet/{bar['id']}?revision={bar['revision']}", headers=ADMIN)
    ov = client.get(f"/api/events/{event}/overview", headers=ADMIN).json()
    stages = next(g["id"] for g in ov["groups"] if g["kind"] == "ethernet" and g["name"] == "Stages")
    client.post(f"/api/events/{event}/records/ethernet", headers=ADMIN, json={
        "venue_id": ov["venues"][0]["id"], "group_id": stages, "data": {"purpose": "Stage 3", "ports": 5}})
    reg = next(r for r in _records(client, event, "ethernet")
               if r["group_id"] != stages and r["id"] not in (office["id"], bar["id"]))
    client.post(f"/api/events/{event}/records/ethernet/{reg['id']}/move", headers=ADMIN,
                json={"revision": reg["revision"], "group_id": stages, "before_id": None})

    draft = _log(client, event)
    assert draft["from_version"] == 2 and draft["counts"] == {"added": 1, "removed": 1, "changed": 3}
    eth = _section(draft, "ethernet")
    assert [i["name"] for i in eth["added"]] == ["Stage 3"] and eth["added"][0]["summary"].startswith("5 points")
    assert [i["name"] for i in eth["removed"]] == ["FoodTrucks / Bar"]
    office_change = next(c for c in eth["changed"] if c["name"] == "Photographers Office")
    assert office_change["fields"] == [{"field": "ports", "label": "Ethernet Points", "old": "2", "new": "4"}]
    moved = next(c for c in eth["changed"] if c["name"] == reg["data"]["purpose"])
    assert moved["fields"][0]["label"] == "Moved" and moved["fields"][0]["new"].endswith("Stages")
    secret = _section(draft, "wifi")["changed"][0]["fields"]
    assert secret == [{"field": "password", "label": "Password value", "old": "", "new": "password changed"}]
    assert "NewPass2!" not in str(draft) and "StaffPass1!" not in str(draft)
    altice = next(p for p in draft["points"] if p["venue"] == "Altice Arena")
    assert (altice["old"], altice["new"]) == (124, 124 - 98 + 2 + 5)

    r = client.post(f"/api/events/{event}/publish", headers=ADMIN, json={"note": "Added stage 3, removed food trucks"})
    assert r.status_code == 201
    v3 = _log(client, event, to=3)
    assert v3["note"] == "Added stage 3, removed food trucks" and v3["counts"] == draft["counts"]
    html = client.app.state.converter.html
    assert "Added stage 3, removed food trucks" in html and "<strong>Removed</strong> FoodTrucks / Bar" in html
    assert "NewPass2!" not in html.split("Changes in this version")[1]

    xlsx = client.app.state.storage.get(version_key(event, 3, "original.xlsx"))
    rows = [[c.value for c in row] for row in load_workbook(io.BytesIO(xlsx))["Changes - Log"].iter_rows()]
    assert rows[0] == ["Change", "Type", "Item", "Where", "Details"]
    assert rows[1][0] == "Note" and ["Removed", "Ethernet requirements", "FoodTrucks / Bar"] in [r[:3] for r in rows]
    assert any(r[0] == "Total" and r[2] == "Altice Arena" for r in rows)


def test_compare_any_two_versions(client: TestClient, event: int) -> None:
    client.post(f"/api/events/{event}/publish", headers=ADMIN)  # V2
    office = _find(client, event, "ethernet", purpose="Photographers Office")
    office = _patch(client, event, office, {"ports": 4})
    client.post(f"/api/events/{event}/publish", headers=ADMIN)  # V3
    _patch(client, event, office, {"ports": 6})
    client.post(f"/api/events/{event}/publish", headers=ADMIN)  # V4
    v4 = _section(_log(client, event, to=4), "ethernet")["changed"][0]["fields"][0]
    assert (v4["old"], v4["new"]) == ("4", "6")
    across = _log(client, event, to=4, **{"from": 2})
    assert across["from_version"] == 2
    assert _section(across, "ethernet")["changed"][0]["fields"][0]["old"] == "2"
    assert _log(client, event)["published_versions"] == [2, 3, 4]

    v1 = _log(client, event, to=1)  # the uploaded workbook: nothing before it
    assert v1["first"] and v1["items"] == 14
    from_upload = _log(client, event, to=4, **{"from": 1})
    fields = next(c for c in _section(from_upload, "ethernet")["changed"] if c["name"] == "Photographers Office")
    assert fields["fields"] == [{"field": "ports", "label": "Ethernet Points", "old": "2", "new": "6"}]
    assert from_upload["counts"] == {"added": 0, "removed": 0, "changed": 1}
    assert client.get(f"/api/events/{event}/changelog", headers=VIEWER, params={"to": 9}).status_code == 404
    assert client.get(f"/api/events/{event}/changelog", headers=VIEWER, params={"to": "x"}).status_code == 422
    assert client.get(f"/api/events/{event}/changelog", headers={}).status_code == 403



def test_edited_workbook_version_is_compared_by_content(client: TestClient, event: int) -> None:
    from openpyxl import load_workbook as load

    client.post(f"/api/events/{event}/publish", headers=ADMIN)  # V2
    # A workbook changed outside the app (as the spreadsheet editor would): one Ethernet item's points.
    wb = load(io.BytesIO(client.app.state.storage.get(version_key(event, 2, "original.xlsx"))))
    ws = next(w for w in wb.worksheets if "Ethernet" in w.title)
    header = next(r for r in ws.iter_rows() if any(c.value == "Purpose" for c in r))
    col = {c.value: c.column for c in header}
    row = next(r for r in ws.iter_rows() if r[col["Purpose"] - 1].value == "Photographers Office")
    ws.cell(row[0].row, col["Ethernet Points"], 9)
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/documents/upload", headers=ADMIN,
                    files={"file": ("LIS26 _ Event_Network_Requirements_V3.xlsx", buf.getvalue(), XLSX)})
    assert r.status_code in (200, 201), r.text

    log = _log(client, event, to=3)  # compared with V2 by default
    assert log["from_version"] == 2 and log["counts"] == {"added": 0, "removed": 0, "changed": 1}
    change = _section(log, "ethernet")["changed"][0]
    assert change["name"] == "Photographers Office" and change["fields"][0]["new"] == "9"
    assert log["notices"] == ["V3 was uploaded, so its items were read from its Excel file."]
    again = _log(client, event, to=3)  # second read comes from the cache
    assert again["counts"] == log["counts"]
