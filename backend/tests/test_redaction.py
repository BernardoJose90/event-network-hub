from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

from app.xlsx import redact_workbook, sensitive_values

from .conftest import xlsx_bytes

PASSWORDS = ["s3cret!", "hunter2", 12345]


def _build(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Devices"
    ws["A1"] = "Network inventory"  # title row above the header
    ws.append([])
    ws.append(["Device", "Username", "  PASSWORD Value ", "VLAN"])  # header in row 3
    ws.append(["sw-core-01", "admin", PASSWORDS[0], 10])
    ws.append(["sw-edge-02", "admin", PASSWORDS[1], 20])
    ws.append(["fw-01", "root", None, 30])  # empty cell must stay empty
    ws.append(["ap-03", "ops", PASSWORDS[2], 40])

    other = wb.create_sheet("Notes")
    other.append(["Topic", "Password"])  # similar but not a sensitive header
    other.append(["wifi", "visible-on-purpose"])


def _redact() -> tuple[Workbook, int]:
    data, count = redact_workbook(xlsx_bytes(_build))
    return load_workbook(io.BytesIO(data)), count


def test_password_column_is_redacted() -> None:
    wb, count = _redact()
    ws = wb["Devices"]
    assert [ws.cell(r, 3).value for r in range(4, 8)] == ["REDACTED", "REDACTED", None, "REDACTED"]
    assert count == 3


def test_no_password_value_survives_anywhere() -> None:
    wb, _ = _redact()
    all_values = {
        cell.value for ws in wb.worksheets for row in ws.iter_rows() for cell in row
    }
    for secret in PASSWORDS:
        assert secret not in all_values


def test_other_columns_and_header_untouched() -> None:
    wb, _ = _redact()
    ws = wb["Devices"]
    assert ws["A1"].value == "Network inventory"
    assert [c.value for c in ws[3]] == ["Device", "Username", "  PASSWORD Value ", "VLAN"]
    assert [ws.cell(r, 1).value for r in range(4, 8)] == ["sw-core-01", "sw-edge-02", "fw-01", "ap-03"]
    assert [ws.cell(r, 2).value for r in range(4, 8)] == ["admin", "admin", "root", "ops"]
    assert [ws.cell(r, 4).value for r in range(4, 8)] == [10, 20, 30, 40]


def test_sheet_without_sensitive_header_untouched() -> None:
    wb, _ = _redact()
    assert [[c.value for c in row] for row in wb["Notes"].iter_rows()] == [
        ["Topic", "Password"],
        ["wifi", "visible-on-purpose"],
    ]


def test_header_below_search_window_is_ignored() -> None:
    def build(wb: Workbook) -> None:
        ws = wb.active
        for _ in range(5):
            ws.append(["filler"])
        ws.append(["Password Value"])  # row 6: outside the first 5 rows
        ws.append(["still-visible"])

    data, count = redact_workbook(xlsx_bytes(build))
    assert count == 0
    assert load_workbook(io.BytesIO(data)).active["A7"].value == "still-visible"


def test_sensitive_values_lists_original_values() -> None:
    assert sensitive_values(xlsx_bytes(_build)) == ["s3cret!", "hunter2", "12345"]
