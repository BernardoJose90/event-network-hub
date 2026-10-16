"""Import an .xlsx workbook into the live sheet model."""

from __future__ import annotations

import io
import re
from collections import Counter
from datetime import date, datetime, time
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..xlsx_colors import resolve, theme_palette
from .model import Cell, Column, Merge, Row, Sheet, Source, Style, Workbook

# Empty rows that only carry a dropdown are kept up to this far past the content
# (the sheets have empty "add a partner here" rows); further ones are dropped.
VALIDATION_ROWS_PAST_CONTENT = 15
DEFAULT_COL_WIDTH_CHARS = 8.43

_RANGE = re.compile(
    r"^(?:'?(?P<sheet>[^'!]+)'?!)?\$?(?P<c1>[A-Z]{1,3})\$?(?P<r1>\d+)(?::\$?(?P<c2>[A-Z]{1,3})\$?(?P<r2>\d+))?$"
)
_SUM = re.compile(r"^=\s*sum\((?P<args>.*)\)\s*$", re.IGNORECASE)
# Written by our exporter for totals with nothing to add up yet (they show 0).
_EMPTY_TOTAL = re.compile(r'^=\s*N\(\s*"(?P<kind>section|grand) total"\s*\)\s*$', re.IGNORECASE)
_A1 = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
_ARITHMETIC = re.compile(r"^=[A-Z0-9$+\-*/(). ]+$")
_WHITE = {"#FFFFFF"}
_BLACK = {"#000000"}


def _col_id(index: int) -> str:
    return f"c{index}"


def _row_id(index: int) -> str:
    return f"r{index}"


def _hex(color: Any, palette: list[str]) -> str | None:
    """'#RRGGBB' for an explicit RGB or theme colour; None for 'no colour'."""
    return resolve(color, palette)


def _font_key(cell) -> tuple[str | None, float | None]:
    font = cell.font
    if font is None:
        return None, None
    return font.name, (float(font.sz) if font.sz else None)


def _style(cell, default_font: str, default_size: float, palette: list[str]) -> Style | None:
    fill = None
    if cell.fill is not None and cell.fill.fill_type == "solid":
        fill = _hex(cell.fill.fgColor, palette)
        if fill in _WHITE:
            fill = None
    font = cell.font
    color = _hex(font.color, palette) if font is not None else None
    if color in _BLACK:
        color = None
    align = cell.alignment.horizontal if cell.alignment is not None else None
    align = {"left": "left", "center": "center", "centerContinuous": "center", "right": "right"}.get(align or "")
    name, size = _font_key(cell)
    style = Style(
        font=name[:60] if name and name != default_font else None,
        bold=True if font is not None and font.b else None,
        italic=True if font is not None and font.i else None,
        size=size if size and size != default_size and 4 <= size <= 72 else None,
        color=color,
        fill=fill,
        align=align,
        wrap=True if cell.alignment is not None and cell.alignment.wrap_text else None,
    )
    return style if style.model_dump(exclude_none=True) else None


def _plain_value(value: Any) -> str | int | float | bool | None:
    if isinstance(value, bool) or value is None:
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)  # Google Sheets stores whole numbers as 1.0
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (date, time)):
        return value.isoformat()
    text = str(value)
    return text if text.strip() else None


def _bounds(ws: Worksheet, palette: list[str]) -> tuple[int, int]:
    last_row = last_col = 0
    for row in ws.iter_rows():
        for cell in row:
            has_value = cell.value is not None and str(cell.value).strip() != ""
            has_fill = cell.fill is not None and cell.fill.fill_type == "solid" and _hex(cell.fill.fgColor, palette) not in (None, *_WHITE)
            if has_value or has_fill:
                last_row = max(last_row, cell.row)
                last_col = max(last_col, cell.column)
    for dv in ws.data_validations.dataValidation:
        for rng in dv.sqref.ranges:
            if rng.min_col <= max(last_col, 1):
                last_row = max(last_row, min(rng.max_row, last_row + VALIDATION_ROWS_PAST_CONTENT))
    return max(last_row, 1), max(last_col, 1)


def _resolve_list_source(wb, ws: Worksheet, formula: str, sheet_ids: dict[str, str]) -> Source | None:
    formula = (formula or "").strip()
    if formula.startswith('"') and formula.endswith('"'):
        values = [v.strip() for v in formula[1:-1].split(",") if v.strip()]
        return Source(values=values) if values else None
    name = wb.defined_names.get(formula) or ws.defined_names.get(formula)
    if name is not None:
        formula = name.attr_text
    m = _RANGE.match(formula.replace(" ", ""))
    if m is None:
        return None
    sheet_name = m.group("sheet") or ws.title
    if sheet_name not in sheet_ids:
        return None
    return Source(sheet=sheet_ids[sheet_name], column=_col_id(column_index_from_string(m.group("c1"))))


def _sum_refs(args: str) -> list[tuple[str, int, int]] | None:
    """Parse SUM arguments into (column letter, first row, last row); None if unsupported."""
    out = []
    for part in [a.strip() for a in args.split(",") if a.strip()]:
        m = _RANGE.match(part.replace("$", ""))
        if m is None or m.group("sheet"):
            return None
        c2 = m.group("c2") or m.group("c1")
        if c2 != m.group("c1"):
            return None
        r1, r2 = int(m.group("r1")), int(m.group("r2") or m.group("r1"))
        out.append((m.group("c1"), min(r1, r2), max(r1, r2)))
    return out


def _classify_formulas(formulas: dict[tuple[int, int], str]) -> dict[tuple[int, int], tuple[str, str | None]]:
    """Map (row, col) -> (kind, expr) for formulas we can represent; others are omitted."""
    kinds: dict[tuple[int, int], tuple[str, str | None]] = {}
    sums: dict[tuple[int, int], list[tuple[str, int, int]] | str] = {}
    for (r, c), f in formulas.items():
        empty = _EMPTY_TOTAL.match(f)
        if empty is not None:
            kinds[(r, c)] = (f"{empty.group('kind').lower()}_total", None)
            continue
        m = _SUM.match(f)
        if m is None:
            continue
        args = m.group("args").strip()
        if args.upper() == "#REF!":
            sums[(r, c)] = "#REF!"
        else:
            refs = _sum_refs(args)
            if refs is not None:
                sums[(r, c)] = refs

    for (r, c), refs in sums.items():
        letter = get_column_letter(c)
        if (r, c) in kinds:
            continue
        if refs == "#REF!":
            kinds[(r, c)] = ("section_total", None)
            continue
        if all(col == letter for col, _, _ in refs) and refs and refs[0][1] == r + 1:
            covered = sorted(row for _, a, b in refs for row in range(a, b + 1))
            if covered == list(range(r + 1, r + 1 + len(covered))):
                kinds[(r, c)] = ("section_total", None)

    for (r, c), refs in sums.items():
        if (r, c) in kinds or refs == "#REF!":
            continue
        letter = get_column_letter(c)
        singles = all(col == letter and a == b for col, a, b in refs)
        # A SUM of single cells in its own column, mostly section totals, is the
        # overall total (one listed header may have no total, e.g. an empty section).
        section_refs = sum(1 for _, a, _ in refs if kinds.get((a, c), ("",))[0] == "section_total")
        if singles and section_refs >= 2:
            kinds[(r, c)] = ("grand_total", None)

    for (r, c), f in formulas.items():
        if (r, c) in kinds or not _ARITHMETIC.match(f.upper()):
            continue
        refs = _A1.findall(f.upper())
        if refs and all(int(row) == r for _, row in refs):
            expr = _A1.sub(lambda m: "{" + _col_id(column_index_from_string(m.group(1))) + "}", f[1:].upper())
            kinds[(r, c)] = ("row_formula", expr.replace(" ", ""))
    return kinds


def import_workbook(data: bytes) -> Workbook:
    wb = load_workbook(io.BytesIO(data))
    cached = load_workbook(io.BytesIO(data), data_only=True)
    palette = theme_palette(wb)
    worksheets = [ws for ws in wb.worksheets if isinstance(ws, Worksheet)]
    sheet_ids = {ws.title: f"s{i + 1}" for i, ws in enumerate(worksheets)}
    warnings: list[str] = []
    sheets: list[Sheet] = []

    for ws in worksheets:
        values_ws = cached[ws.title]
        last_row, last_col = _bounds(ws, palette)
        # The sheet's default font: the most common one among cells with content.
        fonts = Counter(
            _font_key(cell) for row in ws.iter_rows(max_row=last_row, max_col=last_col)
            for cell in row if cell.value is not None
        )
        (default_font, default_size), _ = (fonts.most_common(1) or [((None, None), 0)])[0]
        default_font = (default_font or "Overpass")[:60]
        default_size = default_size if default_size and 4 <= default_size <= 72 else 10.0
        default_width = ws.sheet_format.defaultColWidth or DEFAULT_COL_WIDTH_CHARS
        columns = []
        for c in range(1, last_col + 1):
            dim = ws.column_dimensions.get(get_column_letter(c))
            width = dim.width if dim is not None and dim.width else default_width
            columns.append(Column(id=_col_id(c), width=max(10, min(2000, round(width * 7 + 5)))))

        formulas = {
            (cell.row, cell.column): cell.value
            for row in ws.iter_rows(max_row=last_row, max_col=last_col)
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        }
        kinds = _classify_formulas(formulas)

        rows: list[Row] = []
        for r in range(1, last_row + 1):
            cells: dict[str, Cell] = {}
            for c in range(1, last_col + 1):
                cell = ws.cell(r, c)
                style = _style(cell, default_font, default_size, palette)
                if (r, c) in formulas:
                    kind_expr = kinds.get((r, c))
                    if kind_expr is not None:
                        kind, expr = kind_expr
                        cells[_col_id(c)] = Cell(kind=kind, expr=expr, style=style)
                        continue
                    value = _plain_value(values_ws.cell(r, c).value)
                    warnings.append(f"{ws.title}!{cell.coordinate}: formula {formulas[(r, c)]} kept as its value")
                else:
                    value = _plain_value(cell.value)
                if isinstance(value, bool):
                    cells[_col_id(c)] = Cell(v=value, kind="checkbox", style=style)
                elif value is not None or (style is not None and style.fill):
                    # Empty cells only matter when they are visibly filled.
                    kind = "number" if isinstance(value, (int, float)) else "text"
                    cells[_col_id(c)] = Cell(v=value, kind=kind, style=style)
            rows.append(Row(id=_row_id(r), cells=cells))

        for dv in ws.data_validations.dataValidation:
            if dv.type != "list":
                continue
            source = _resolve_list_source(wb, ws, dv.formula1, sheet_ids)
            if source is None:
                warnings.append(f"{ws.title}: dropdown source {dv.formula1!r} not supported; cells kept as text")
                continue
            for rng in dv.sqref.ranges:
                for r in range(rng.min_row, min(rng.max_row, last_row) + 1):
                    for c in range(rng.min_col, min(rng.max_col, last_col) + 1):
                        cells = rows[r - 1].cells
                        existing = cells.get(_col_id(c))
                        if existing is not None and existing.kind not in ("text", "number"):
                            continue
                        cells[_col_id(c)] = Cell(
                            v=None if existing is None else existing.v,
                            kind="dropdown",
                            source=source,
                            style=None if existing is None else existing.style,
                        )

        merges = []
        for mr in ws.merged_cells.ranges:
            if mr.min_row > last_row or mr.min_col > last_col:
                continue
            if mr.min_row != mr.max_row:
                warnings.append(f"{ws.title}!{mr.coord}: merges across rows are shown unmerged")
                continue
            span = min(mr.max_col, last_col) - mr.min_col + 1
            if span >= 2:
                merges.append(Merge(row=_row_id(mr.min_row), col=_col_id(mr.min_col), span=span))

        # Cells hidden under a merge carry no content of their own.
        col_order = [c.id for c in columns]
        for m in merges:
            start = col_order.index(m.col)
            cells = rows[int(m.row[1:]) - 1].cells
            for covered in col_order[start + 1:start + m.span]:
                cells.pop(covered, None)

        frozen = 0
        if ws.freeze_panes:
            m = re.match(r"^[A-Z]+(\d+)$", ws.freeze_panes)
            frozen = max(0, min(20, int(m.group(1)) - 1)) if m else 0

        sheets.append(
            Sheet(id=sheet_ids[ws.title], name=ws.title[:100], frozen_rows=frozen,
                  font=default_font, font_size=default_size,
                  columns=columns, rows=rows, merges=merges)
        )

    # Dropdowns whose source column was empty (and trimmed away) become plain text.
    columns = {s.id: {c.id for c in s.columns} for s in sheets}
    names = {s.id: s.name for s in sheets}
    for sheet in sheets:
        for row in sheet.rows:
            for col_id, cell in row.cells.items():
                src = cell.source
                if src is not None and src.sheet is not None and src.column not in columns.get(src.sheet, ()):
                    row.cells[col_id] = Cell(v=cell.v, kind="number" if isinstance(cell.v, (int, float)) and not isinstance(cell.v, bool) else "text", style=cell.style)
                    warnings.append(f"{sheet.name}: a dropdown points at an empty column of {names.get(src.sheet, 'another sheet')}; kept as text")

    return Workbook(sheets=sheets, warnings=warnings[:500])
