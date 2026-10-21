"""Post-event reporting: KPIs, CSV/Excel export, recurring issues.

KPIs follow the project brief:
- time to acknowledge: share acknowledged within the SLA (5 minutes), and the median;
- coverage: tickets / (tickets + issues handled outside the system), when that count is given;
- closure quality: resolved/closed tickets with a documented root cause and resolution;
- partner-impacting: count and mean time to resolve;
- recurrence: share of tickets matching an earlier ticket's category and area (any event).
Duplicates merged into another ticket are counted separately and left out of the KPIs.
"""

from __future__ import annotations

import csv
import io
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from statistics import mean, median
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..db import Connection
from . import service, store
from .models import CATEGORIES, DONE_STATUSES, SEVERITIES, SOURCES, STATUSES, TEAMS


def _area(ticket: dict[str, Any]) -> str | None:
    """Where the issue was: stand, else location; normalised for matching."""
    value = ticket.get("stand") or ticket.get("location")
    return re.sub(r"\s+", " ", value).strip().lower() if value else None


def _pct(part: int, whole: int) -> float | None:
    return round(100 * part / whole, 1) if whole else None


def _avg(values: list[float]) -> float | None:
    return round(mean(values), 1) if values else None


def _med(values: list[float]) -> float | None:
    return round(median(values), 1) if values else None


def is_partner_impacting(ticket: dict[str, Any]) -> bool:
    return ticket["category"] == "partner_stand" or bool(ticket.get("partner"))


def kpis(conn: Connection, event_id: int | None, informal: int | None = None) -> dict[str, Any]:
    now = datetime.now(timezone.utc)
    everything = store.list_tickets(conn, limit=1_000_000)
    scoped = [t for t in everything if event_id is None or t["event_id"] == event_id]
    duplicates = [t for t in scoped if t["merged_into"]]
    tickets = [t for t in scoped if not t["merged_into"]]
    slas = {t["id"]: service.sla(t, now) for t in tickets}

    ack_decided = [t for t in tickets if slas[t["id"]]["ack"] != "pending"]
    ack_met = [t for t in ack_decided if slas[t["id"]]["ack"] == "met"]
    ack_minutes = [slas[t["id"]]["minutes_to_ack"] for t in tickets if slas[t["id"]]["minutes_to_ack"] is not None]

    done = [t for t in tickets if t["status"] in DONE_STATUSES]
    documented = [t for t in done if (t["root_cause"] or "").strip() and (t["resolution"] or "").strip()]
    resolve_minutes = [slas[t["id"]]["minutes_to_resolve"] for t in done if slas[t["id"]]["minutes_to_resolve"] is not None]

    partner = [t for t in tickets if is_partner_impacting(t)]
    partner_done = [slas[t["id"]]["minutes_to_resolve"] for t in partner
                    if t["status"] in DONE_STATUSES and slas[t["id"]]["minutes_to_resolve"] is not None]

    # Recurrence: an earlier ticket (any event) with the same category and area.
    seen: set[tuple[str, str]] = set()
    recurring_ids: set[int] = set()
    for t in sorted((t for t in everything if not t["merged_into"]), key=lambda t: t["created_at"]):
        area = _area(t)
        if not area:
            continue
        key = (t["category"], area)
        if key in seen:
            recurring_ids.add(t["id"])
        seen.add(key)
    with_area = [t for t in tickets if _area(t)]
    recurring = [t for t in with_area if t["id"] in recurring_ids]

    by_severity = {}
    for sev in SEVERITIES:
        group = [t for t in tickets if t["severity"] == sev]
        decided = [t for t in group if slas[t["id"]]["resolve"] in ("met", "breached")]
        by_severity[sev] = {
            "count": len(group),
            "resolved_within_sla_pct": _pct(sum(slas[t["id"]]["resolve"] == "met" for t in decided), len(decided)),
            "mean_minutes_to_resolve": _avg([slas[t["id"]]["minutes_to_resolve"] for t in group
                                             if slas[t["id"]]["minutes_to_resolve"] is not None]),
        }

    return {
        "event_id": event_id,
        "tickets": len(tickets),
        "duplicates": len(duplicates),
        "open": sum(t["status"] not in DONE_STATUSES for t in tickets),
        "by_status": dict(Counter(t["status"] for t in tickets)),
        "by_category": dict(Counter(t["category"] for t in tickets)),
        "by_team": dict(Counter(t["team"] for t in tickets)),
        "by_source": dict(Counter(t["source"] for t in tickets)),
        "by_severity": by_severity,
        "acknowledge": {
            "target_minutes": SEVERITIES["P1"]["ack_minutes"],
            "within_target_pct": _pct(len(ack_met), len(ack_decided)),
            "goal_pct": 90,
            "median_minutes": _med(ack_minutes),
        },
        "coverage": {
            "informal_issues": informal,
            "pct": _pct(len(tickets), len(tickets) + informal) if informal is not None else None,
            "goal_pct": 90,
        },
        "closure_quality": {"closed": len(done), "documented": len(documented),
                            "pct": _pct(len(documented), len(done)), "goal_pct": 90},
        "resolve": {"mean_minutes": _avg(resolve_minutes), "median_minutes": _med(resolve_minutes)},
        "partner_impacting": {"count": len(partner), "resolved": len(partner_done),
                              "mean_minutes_to_resolve": _avg(partner_done)},
        "recurrence": {"with_area": len(with_area), "recurring": len(recurring),
                       "pct": _pct(len(recurring), len(with_area))},
        "top_areas": [{"area": a, "count": n} for a, n in Counter(
            (t.get("stand") or t.get("location")) for t in tickets if _area(t)).most_common(10)],
    }


def recurring_issues(conn: Connection, min_count: int = 2) -> list[dict[str, Any]]:
    """Issues seen at least min_count times in the same category and area, across events,
    with what fixed them: the knowledge base."""
    groups: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for t in store.list_tickets(conn, limit=1_000_000):
        if t["merged_into"]:
            continue
        area = _area(t)
        if area:
            groups[(t["category"], area)].append(t)
    out = []
    for (category, _), items in groups.items():
        if len(items) < min_count:
            continue
        items.sort(key=lambda t: t["created_at"], reverse=True)
        fixes = [{"ticket": t["id"], "event": t["event_code"], "root_cause": t["root_cause"],
                  "resolution": t["resolution"], "at": t["resolved_at"] or t["closed_at"]}
                 for t in items if t["resolution"]]
        out.append({
            "category": category,
            "area": items[0].get("stand") or items[0].get("location"),
            "count": len(items),
            "events": sorted({t["event_code"] for t in items if t["event_code"]}),
            "last_seen": items[0]["created_at"],
            "open": sum(t["status"] not in DONE_STATUSES for t in items),
            "fixes": fixes[:5],
        })
    out.sort(key=lambda g: g["last_seen"], reverse=True)
    out.sort(key=lambda g: g["count"], reverse=True)  # stable: most frequent, then most recent
    return out


# --- export -------------------------------------------------------------------------------

COLUMNS = [
    ("Ticket", 8), ("Event", 10), ("Created (UTC)", 18), ("Title", 40), ("Description", 50), ("Category", 22),
    ("Severity", 11), ("Status", 13), ("Team", 14), ("Owner", 18), ("Location", 22), ("Stand", 12),
    ("Partner", 22), ("Reporter", 18), ("Reporter phone", 16), ("Source", 14), ("Acknowledged (UTC)", 18),
    ("Minutes to acknowledge", 12), ("Acknowledged within SLA", 12), ("Resolved (UTC)", 18),
    ("Minutes to resolve", 12), ("Resolved within SLA", 12), ("Root cause", 40), ("Resolution", 40),
    ("Duplicate of", 10),
]

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


_PLAIN_NUMBER = re.compile(r"[+-]?[\d\s().-]+")  # phone numbers and numbers can't run formulas


def _safe(value: Any) -> Any:
    """Neutralise spreadsheet formula injection: reporter text comes from WhatsApp."""
    if isinstance(value, str) and value.startswith(_FORMULA_START) and not _PLAIN_NUMBER.fullmatch(value):
        return "'" + value
    return value


def _when(value: str | None) -> str | None:
    return value.replace("T", " ")[:16] if value else None


def export_rows(conn: Connection, event_id: int | None) -> list[list[Any]]:
    now = datetime.now(timezone.utc)
    rows = []
    for t in sorted(store.list_tickets(conn, event_id=event_id, limit=1_000_000), key=lambda t: t["id"]):
        s = service.sla(t, now)
        yes_no = {"met": "Yes", "breached": "No"}
        rows.append([_safe(v) for v in (
            t["id"], t["event_code"], _when(t["created_at"]), t["title"], t["description"],
            CATEGORIES[t["category"]], SEVERITIES[t["severity"]]["label"], STATUSES[t["status"]], TEAMS[t["team"]],
            t["assignee"], t["location"], t["stand"], t["partner"], t["reporter_name"], t["reporter_phone"],
            SOURCES[t["source"]], _when(t["acknowledged_at"]), s["minutes_to_ack"], yes_no.get(s["ack"]),
            _when(t["resolved_at"] or t["closed_at"]), s["minutes_to_resolve"], yes_no.get(s["resolve"]),
            t["root_cause"], t["resolution"], t["merged_into"],
        )])
    return rows


def export_csv(conn: Connection, event_id: int | None) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([c for c, _ in COLUMNS])
    writer.writerows(export_rows(conn, event_id))
    return buffer.getvalue().encode("utf-8-sig")  # BOM so Excel opens it as UTF-8


def export_xlsx(conn: Connection, event_id: int | None, title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"
    header_fill = PatternFill("solid", fgColor="0C1439")
    for col, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, col, name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(col)].width = width
    for r, row in enumerate(export_rows(conn, event_id), start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(r, c, value).alignment = Alignment(wrap_text=c in (4, 5, 23, 24), vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.properties.title = title
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
