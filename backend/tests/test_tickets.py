from __future__ import annotations

import csv
import hashlib
import hmac
import io
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import db
from app.config import Settings
from app.main import create_app
from app.tickets import outbox, service, store

ADMIN = {"X-Role": "network-admin", "X-User": "alice"}
VIEWER = {"X-Role": "viewer", "X-User": "victor"}
BOT = {"Authorization": "Bearer bot-secret-token"}


def _settings(tmp_path: Path, **extra) -> Settings:
    return Settings(data_dir=tmp_path, gotenberg_url="http://unused", gotenberg_timeout_seconds=1,
                    max_upload_bytes=1024 * 1024, **extra)


@pytest.fixture
def client(tmp_path: Path):
    settings = _settings(tmp_path, bot_api_token="bot-secret-token")
    app = create_app(settings)
    with TestClient(app) as c:
        c.settings = settings
        yield c


def _event(client, code="LIS26") -> int:
    return client.post("/api/events", headers=ADMIN, json={"event_code": code}).json()["id"]


def _ticket(client, **fields) -> dict:
    body = {"title": "Wi-Fi down in Hall 1", "category": "wifi", "severity": "P2"} | fields
    r = client.post("/api/tickets", headers=ADMIN, json=body)
    assert r.status_code == 201, r.text
    return r.json()


def _backdate(client, ticket_id: int, minutes: float, **columns) -> None:
    """Move a ticket's created time (and given timestamp columns) back by minutes."""
    conn = db.connect(client.settings.db_path)
    created = datetime.now(timezone.utc) - timedelta(minutes=minutes)
    values = {"created_at": created.isoformat(timespec="seconds")}
    for column, offset in columns.items():
        values[column] = (created + timedelta(minutes=offset)).isoformat(timespec="seconds")
    with conn:
        conn.execute(f"UPDATE tickets SET {', '.join(f'{k} = ?' for k in values)} WHERE id = ?",
                     (*values.values(), ticket_id))
    conn.close()


# --- create, roles -------------------------------------------------------------------------


def test_raise_ticket_and_roles(client: TestClient) -> None:
    event = _event(client)
    t = _ticket(client, event_id=event, location="Hall 1", category="hardware")
    assert (t["status"], t["team"], t["source"], t["event_code"]) == ("new", "it_support", "manual", "LIS26")
    assert t["sla"]["ack"] == "pending" and t["sla"]["resolve"] == "pending"
    assert client.post("/api/tickets", headers=VIEWER, json={"title": "x", "category": "wifi"}).status_code == 403
    assert client.get("/api/tickets", headers=VIEWER).json()[0]["id"] == t["id"]  # everyone on shift can see
    assert client.get("/api/tickets", headers={}).status_code == 403
    detail = client.get(f"/api/tickets/{t['id']}", headers=VIEWER).json()
    assert detail["activity"][0]["kind"] == "created" and "IT Support" not in detail["activity"][0]["body"]
    assert client.post("/api/tickets", headers=ADMIN, json={"title": "", "category": "wifi"}).status_code == 422
    assert client.post("/api/tickets", headers=ADMIN, json={"title": "x", "category": "nope"}).status_code == 422


def test_lifecycle_acknowledge_resolve_reopen_close(client: TestClient) -> None:
    t = _ticket(client)
    url = f"/api/tickets/{t['id']}"
    r = client.post(f"{url}/acknowledge", headers=ADMIN, json={"revision": t["revision"]})
    t = r.json()
    assert (t["status"], t["assignee"], t["acknowledged_by"]) == ("acknowledged", "alice", "alice")
    assert t["sla"]["ack"] == "met"

    r = client.post(f"{url}/status", headers=ADMIN, json={"revision": t["revision"], "status": "resolved"})
    assert r.status_code == 422 and "root cause" in r.json()["detail"]
    r = client.post(f"{url}/status", headers=ADMIN, json={"revision": t["revision"], "status": "resolved",
                                                          "root_cause": "AP lost PoE", "resolution": "Reseated cable"})
    t = r.json()
    assert t["status"] == "resolved" and t["resolved_at"] and t["sla"]["resolve"] == "met"

    t = client.post(f"{url}/status", headers=ADMIN, json={"revision": t["revision"], "status": "in_progress"}).json()
    assert t["status"] == "in_progress" and t["resolved_at"] is None  # reopened
    r = client.post(f"{url}/status", headers=ADMIN, json={"revision": t["revision"], "status": "new"})
    assert r.status_code == 422

    t = client.post(f"{url}/status", headers=ADMIN, json={"revision": t["revision"], "status": "closed"}).json()
    assert t["status"] == "closed" and t["closed_at"] and t["root_cause"] == "AP lost PoE"  # kept from before

    kinds = [(a["kind"], a["new"]) for a in client.get(url, headers=ADMIN).json()["activity"]]
    assert kinds == [("created", None), ("acknowledge", "acknowledged"), ("assign", "alice"),
                     ("status", "resolved"), ("status", "in_progress"), ("status", "closed")]


def test_moving_on_counts_as_acknowledging(client: TestClient) -> None:
    t = _ticket(client)
    t = client.post(f"/api/tickets/{t['id']}/status", headers=ADMIN,
                    json={"revision": t["revision"], "status": "in_progress"}).json()
    assert t["acknowledged_at"] and t["acknowledged_by"] == "alice"


def test_edit_assign_and_stale_revision(client: TestClient) -> None:
    t = _ticket(client)
    url = f"/api/tickets/{t['id']}"
    r = client.patch(url, headers=ADMIN, json={"revision": t["revision"],
                                               "changes": {"severity": "P1", "partner": "IBM", "stand": "A12"}})
    assert r.status_code == 200 and r.json()["severity"] == "P1"
    assert client.patch(url, headers=ADMIN, json={"revision": t["revision"], "changes": {"title": "x"}}).status_code == 409
    t = r.json()
    t = client.post(f"{url}/assign", headers=ADMIN, json={"revision": t["revision"], "assignee": "bob",
                                                          "team": "it_support"}).json()
    assert (t["assignee"], t["team"]) == ("bob", "it_support")
    assert "bob" in client.get("/api/tickets/meta", headers=VIEWER).json()["assignees"]
    fields = [a["field"] for a in client.get(url, headers=ADMIN).json()["activity"] if a["field"]]
    assert set(fields) >= {"severity", "partner", "stand", "assignee", "team"}
    assert client.patch(url, headers=VIEWER, json={"revision": t["revision"], "changes": {}}).status_code == 403


def test_filters_and_search(client: TestClient) -> None:
    a = _ticket(client, title="Stand A12 offline", partner="IBM", category="partner_stand", location="Hall 2")
    b = _ticket(client, title="Slow Wi-Fi", category="wifi", severity="P3")
    c = _ticket(client, title="Switch fault", category="hardware")
    client.post(f"/api/tickets/{c['id']}/status", headers=ADMIN,
                json={"revision": c["revision"], "status": "closed", "root_cause": "x", "resolution": "y"})
    ids = lambda **q: [t["id"] for t in client.get("/api/tickets", headers=VIEWER, params=q).json()]
    assert ids() == [b["id"], a["id"]]  # open by default, newest first
    assert ids(status="done") == [c["id"]]
    assert ids(status="all") == [c["id"], b["id"], a["id"]]
    assert ids(partner="ibm") == [a["id"]]
    assert ids(q="slow") == [b["id"]]
    assert ids(q=f"#{a['id']}") == [a["id"]]
    assert ids(severity="P3") == [b["id"]]
    assert ids(unassigned="true") == [b["id"], a["id"]]
    assert ids(q="100%_") == []  # LIKE wildcards are literal
    assert client.get("/api/tickets", headers=VIEWER, params={"status": "bogus"}).status_code == 422


def test_comments(client: TestClient) -> None:
    t = _ticket(client)
    r = client.post(f"/api/tickets/{t['id']}/comments", headers=ADMIN, json={"body": "Checking the switch"})
    assert r.status_code == 201
    r = client.post(f"/api/tickets/{t['id']}/comments", headers=ADMIN,
                    json={"body": "On our way", "visibility": "reporter"})
    assert r.status_code == 422  # manual tickets have no WhatsApp reporter
    assert client.get(f"/api/tickets/{t['id']}", headers=ADMIN).json()["activity"][-1]["body"] == "Checking the switch"


# --- WhatsApp bot -------------------------------------------------------------------------


def test_bot_creates_tickets_idempotently(client: TestClient) -> None:
    event = _event(client, "WEB26")
    report = {"external_id": "wa-123", "message": "Stand B4 has no internet\nPartner says since 10:00",
              "event_code": "web26", "partner": "Acme", "stand": "B4", "reporter_name": "Ana",
              "reporter_phone": "+351900000000"}
    assert client.post("/api/bot/tickets", json=report).status_code == 401
    assert client.post("/api/bot/tickets", headers={"Authorization": "Bearer nope"}, json=report).status_code == 401
    r = client.post("/api/bot/tickets", headers=BOT, json=report)
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["created"] and body["status"] == "new" and body["url"].endswith(f"?ticket={body['ticket']}")
    again = client.post("/api/bot/tickets", headers=BOT, json=report)
    assert again.status_code == 200 and again.json()["ticket"] == body["ticket"] and not again.json()["created"]

    t = client.get(f"/api/tickets/{body['ticket']}", headers=ADMIN).json()
    assert (t["source"], t["category"], t["event_id"], t["title"]) == ("whatsapp", "partner_stand", event,
                                                                         "Stand B4 has no internet")
    assert client.get("/api/bot/tickets/wa-123", headers=BOT).json()["ticket"] == body["ticket"]
    assert client.get("/api/bot/tickets/unknown", headers=BOT).status_code == 404


def test_bot_reports_go_to_the_live_event(client: TestClient) -> None:
    event = _event(client)
    assert client.put("/api/helpdesk/live-event", headers=VIEWER, json={"event_id": event}).status_code == 403
    assert client.put("/api/helpdesk/live-event", headers=ADMIN, json={"event_id": event}).status_code == 200
    r = client.post("/api/bot/tickets", headers=BOT, json={"external_id": "wa-1", "message": "No Wi-Fi in Hall 3"})
    t = client.get(f"/api/tickets/{r.json()['ticket']}", headers=ADMIN).json()
    assert (t["event_id"], t["category"], t["severity"]) == (event, "uncategorised", "P3")
    assert client.get("/api/tickets/meta", headers=VIEWER).json()["live_event_id"] == event


def test_bot_disabled_without_token(tmp_path: Path) -> None:
    with TestClient(create_app(_settings(tmp_path))) as c:
        r = c.post("/api/bot/tickets", headers=BOT, json={"external_id": "x", "message": "y"})
        assert r.status_code == 503


# --- merge, reporter updates ------------------------------------------------------------------


def _bot_ticket(conn, external_id, notify=True, **extra):
    report = {"external_id": external_id, "message": "Wi-Fi down", "reporter_phone": "+1", **extra}
    return service.create_from_bot(conn, report, notify=notify)[0]


def test_merge_and_reporter_updates(tmp_path: Path) -> None:
    db.init_db(tmp_path / "db.sqlite3")
    conn = db.connect(tmp_path / "db.sqlite3")
    main = _bot_ticket(conn, "wa-1", reporter_phone="+111")
    dup = _bot_ticket(conn, "wa-2", reporter_phone="+222")
    merged = service.merge(conn, dup["id"], revision=dup["revision"], into=main["id"], user="alice", notify=True)
    assert (merged["status"], merged["merged_into"], merged["resolution"]) == ("closed", main["id"], f"Duplicate of #{main['id']}")
    with pytest.raises(service.TicketError) as err:
        service.change_status(conn, dup["id"], revision=merged["revision"], status="in_progress", user="a", notify=True)
    assert err.value.status == 409
    with pytest.raises(service.TicketError):
        service.merge(conn, main["id"], revision=store.get_ticket(conn, main["id"])["revision"], into=dup["id"],
                      user="a", notify=True)

    main = store.get_ticket(conn, main["id"])
    service.acknowledge(conn, main["id"], revision=main["revision"], user="alice", notify=True)
    payloads = [json.loads(r["payload"]) for r in conn.execute("SELECT payload FROM ticket_outbox ORDER BY id")]
    # created x2, merge notice to the duplicate's reporter, then the acknowledge goes to both reporters
    assert [p["reporter_phone"] for p in payloads] == ["+111", "+222", "+222", "+111", "+222"]
    assert all(p["ticket"] == main["id"] for p in payloads[-2:])
    assert "duplicate" in payloads[2]["message"]
    conn.close()


def test_no_updates_queued_when_bot_callbacks_are_off(tmp_path: Path) -> None:
    db.init_db(tmp_path / "db.sqlite3")
    conn = db.connect(tmp_path / "db.sqlite3")
    _bot_ticket(conn, "wa-1", notify=False)
    assert conn.execute("SELECT COUNT(*) FROM ticket_outbox").fetchone()[0] == 0
    conn.close()


def test_outbox_delivery_signs_and_retries(tmp_path: Path) -> None:
    path = tmp_path / "db.sqlite3"
    db.init_db(path)
    conn = db.connect(path)
    _bot_ticket(conn, "wa-1")
    received = []

    def ok(request: httpx.Request) -> httpx.Response:
        received.append(request)
        return httpx.Response(200)

    with httpx.Client(transport=httpx.MockTransport(ok)) as client:
        assert outbox.deliver_due(path, "https://bot.example/updates", "s3cret", client) == 1
    body = received[0].content
    assert received[0].headers["X-Signature"] == "sha256=" + hmac.new(b"s3cret", body, hashlib.sha256).hexdigest()
    assert json.loads(body)["external_id"] == "wa-1"
    assert conn.execute("SELECT status FROM ticket_outbox").fetchone()[0] == "sent"

    _bot_ticket(conn, "wa-2")
    with httpx.Client(transport=httpx.MockTransport(lambda r: httpx.Response(500))) as client:
        assert outbox.deliver_due(path, "https://bot.example/updates", "s3cret", client) == 0
        row = conn.execute("SELECT status, attempts, next_try_at, last_error FROM ticket_outbox WHERE id = 2").fetchone()
        assert (row["status"], row["attempts"]) == ("pending", 1) and "500" in row["last_error"]
        assert outbox.deliver_due(path, "https://bot.example/updates", "s3cret", client) == 0  # not due yet
        assert conn.execute("SELECT attempts FROM ticket_outbox WHERE id = 2").fetchone()[0] == 1
        for _ in range(outbox.MAX_ATTEMPTS):
            with conn:
                conn.execute("UPDATE ticket_outbox SET next_try_at = '2000-01-01T00:00:00+00:00' WHERE id = 2")
            outbox.deliver_due(path, "https://bot.example/updates", "s3cret", client)
    assert conn.execute("SELECT status FROM ticket_outbox WHERE id = 2").fetchone()[0] == "failed"
    conn.close()


def test_callback_url_needs_a_secret(tmp_path: Path) -> None:
    app = create_app(_settings(tmp_path, bot_callback_url="https://bot.example/updates"))
    with pytest.raises(RuntimeError, match="BOT_CALLBACK_SECRET"):
        with TestClient(app):
            pass


# --- SLA, reports, export ------------------------------------------------------------------------


def test_sla_states(client: TestClient) -> None:
    t = _ticket(client, severity="P1")
    _backdate(client, t["id"], 20)  # 20 min old, never acknowledged: both P1 targets missed
    t = client.get(f"/api/tickets/{t['id']}", headers=ADMIN).json()
    assert (t["sla"]["ack"], t["sla"]["resolve"]) == ("breached", "breached")
    u = _ticket(client, severity="P4")
    _backdate(client, u["id"], 10, acknowledged_at=3)
    u = client.get(f"/api/tickets/{u['id']}", headers=ADMIN).json()
    assert (u["sla"]["ack"], u["sla"]["resolve"], u["sla"]["minutes_to_ack"]) == ("met", "pending", 3.0)


def test_report_kpis(client: TestClient) -> None:
    lis, van = _event(client, "LIS26"), _event(client, "VAN26")
    # VAN: an earlier ticket in Hall 1, so LIS's Hall 1 Wi-Fi ticket is a recurrence
    old = _ticket(client, event_id=van, location="Hall 1")
    _backdate(client, old["id"], 60 * 24 * 30)
    fast = _ticket(client, event_id=lis, location="hall 1 ", partner=None)
    slow = _ticket(client, event_id=lis, category="partner_stand", partner="IBM", stand="A12", severity="P2")
    open_ = _ticket(client, event_id=lis, category="wired", location="Hall 9")
    dup = _ticket(client, event_id=lis)
    for t, ack, res in ((fast, 2, 10), (slow, 8, 45)):
        r = client.post(f"/api/tickets/{t['id']}/status", headers=ADMIN, json={
            "revision": t["revision"], "status": "resolved", "root_cause": "cause", "resolution": "fix"})
        assert r.status_code == 200
        _backdate(client, t["id"], 60, acknowledged_at=ack, resolved_at=res)
    _backdate(client, open_["id"], 1)
    client.post(f"/api/tickets/{dup['id']}/merge", headers=ADMIN, json={"revision": dup["revision"], "into": fast["id"]})

    k = client.get("/api/tickets/report", headers=VIEWER, params={"event_id": lis, "informal": 1}).json()
    assert (k["event"], k["tickets"], k["duplicates"], k["open"]) == ("LIS26", 3, 1, 1)
    assert k["acknowledge"]["within_target_pct"] == 50.0  # fast met, slow breached, open still pending
    assert k["acknowledge"]["median_minutes"] == 5.0
    assert k["coverage"]["pct"] == 75.0
    assert k["closure_quality"] == {"closed": 2, "documented": 2, "pct": 100.0, "goal_pct": 90}
    assert k["partner_impacting"] == {"count": 1, "resolved": 1, "mean_minutes_to_resolve": 45.0}
    assert k["recurrence"] == {"with_area": 3, "recurring": 1, "pct": 33.3}
    assert k["by_severity"]["P2"]["resolved_within_sla_pct"] == 50.0  # fast 10 min met, slow 45 > 30 min missed
    everything = client.get("/api/tickets/report", headers=VIEWER).json()
    assert (everything["event"], everything["tickets"]) == ("All events", 4)

    recurring = client.get("/api/tickets/recurring", headers=VIEWER).json()
    assert len(recurring) == 1 and recurring[0]["count"] == 2 and recurring[0]["events"] == ["LIS26", "VAN26"]
    assert recurring[0]["fixes"][0]["resolution"] == "fix"


def test_export_is_safe_and_complete(client: TestClient) -> None:
    event = _event(client)
    _ticket(client, event_id=event, title="=HYPERLINK(\"http://evil\")", description="+cmd")
    r = client.get("/api/tickets/export", headers=VIEWER, params={"event_id": event, "format": "csv"})
    assert r.status_code == 200 and 'filename="ICT_tickets_LIS26.csv"' in r.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0][0] == "Ticket" and rows[1][3].startswith("'=") and rows[1][4] == "'+cmd"
    r = client.get("/api/tickets/export", headers=VIEWER, params={"format": "xlsx"})
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A1"].value == "Ticket" and ws["D2"].value.startswith("'=") and ws.max_row == 2
    assert client.get("/api/tickets/export", headers=VIEWER, params={"format": "pdf"}).status_code == 422


def test_tickets_outlive_their_event(client: TestClient) -> None:
    event = _event(client)
    t = _ticket(client, event_id=event)
    client.post(f"/api/documents/{event}/archive", headers=ADMIN)
    assert client.delete(f"/api/documents/{event}", headers=ADMIN).status_code == 204
    t = client.get(f"/api/tickets/{t['id']}", headers=ADMIN).json()
    assert (t["event_id"], t["event_code"]) == (None, "LIS26")
