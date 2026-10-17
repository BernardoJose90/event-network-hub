from __future__ import annotations

import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import db
from app.config import Settings
from app.main import create_app

from .conftest import FakeConverter
from .records_fixture import records_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ADMIN = {"X-Role": "network-admin", "X-User": "alice"}
VIEWER = {"X-Role": "viewer", "X-User": "victor"}


@pytest.fixture
def client(tmp_path: Path):
    settings = Settings(data_dir=tmp_path, gotenberg_url="http://unused", gotenberg_timeout_seconds=1,
                        max_upload_bytes=1024 * 1024)
    app = create_app(settings)
    with TestClient(app) as c:
        app.state.converter = FakeConverter()
        c.settings = settings
        yield c


@pytest.fixture
def event(client: TestClient) -> int:
    r = client.post("/api/documents/upload", headers=ADMIN,
                    files={"file": ("LIS26 _ Event_Network_Requirements_V1.xlsx", records_workbook(), XLSX)})
    doc = r.json()["document"]["id"]
    assert client.post(f"/api/events/{doc}/import", headers=ADMIN, json={}).status_code == 200
    return doc


def _report(client, **params):
    r = client.get("/api/reports", headers=VIEWER, params=params)
    assert r.status_code == 200, r.text
    return r.json()


def test_report_for_an_event(client: TestClient, event: int) -> None:
    client.post(f"/api/events/{event}/publish", headers=ADMIN)  # V2, first published
    ov = client.get(f"/api/events/{event}/overview", headers=ADMIN).json()
    location = next(i["value"] for i in ov["lists"] if i["kind"] == "location")
    venue = ov["venues"][0]["id"]
    section = next(g["id"] for g in ov["groups"] if g["kind"] == "ethernet")

    # ICT requests: one approved on time (4 points), one late and unsure
    token = client.post(f"/api/events/{event}/ict-requests/links", headers=ADMIN, json={"team": "av"}).json()["token"]
    req = client.post(f"/api/request-links/{token}/requests", json={
        "requester_name": "Ana", "requester_contact": "ana@x.com",
        "lines": [{"kind": "wired", "area": location, "purpose": "Encoders", "ports": 4}]}).json()["requests"][0]
    client.post(f"/api/events/{event}/ict-requests/{req['id']}/approve", headers=ADMIN,
                json={"revision": req["revision"], "venue_id": venue, "group_id": section})
    client.put(f"/api/events/{event}/ict-requests/settings", headers=ADMIN, json={"deadline": "2000-01-01"})
    client.post(f"/api/request-links/{token}/requests", json={
        "requester_name": "Ana", "requester_contact": "ana@x.com",
        "lines": [{"kind": "wired", "area": location, "purpose": "Spare", "ports": None}]})

    # help desk: a wired ticket, acknowledged straight away
    t = client.post("/api/tickets", headers=ADMIN, json={"title": "No link", "category": "wired", "event_id": event}).json()
    client.post(f"/api/tickets/{t['id']}/acknowledge", headers=ADMIN, json={"revision": t["revision"]})

    client.post(f"/api/events/{event}/publish", headers=ADMIN)  # V3

    report = _report(client)
    assert [e["code"] for e in report["events"]] == ["LIS26"] and report["goals"]["ack_within_target_pct"] == 90
    r = report["events"][0]
    cp = r["cable_planning"]
    assert cp["points"] == 126 + 4 and cp["points_via_requests"] == 4 and cp["via_requests_pct"] == round(400 / 130, 1)
    assert cp["deadline"] == "2000-01-01" and cp["items_added_after_deadline"] >= 1 and cp["wired_tickets_on_site"] == 1
    rq = r["requests"]
    assert (rq["requests"], rq["teams_using"], rq["on_time_pct"], rq["not_sure_pct"]) == (2, 1, 50.0, 100.0)
    assert rq["by_team"]["av"]["approved"] == 1 and rq["by_team"]["av"]["ports_requested"] == 4 and rq["open"] == 1
    assert "register" not in r  # the dependency register was removed
    ops = r["operations"]
    assert (ops["tickets"], ops["ack_within_target_pct"]) == (1, 100.0)
    pub = r["publishing"]
    assert (pub["versions"], pub["published"], pub["unpublished_changes"]) == (3, 2, 0)
    assert pub["first_to_last"]["from"] == 2 and pub["first_to_last"]["to"] == 3
    assert pub["first_to_last"]["added"] == 1 and pub["first_to_last"]["points_delta"] == 4


def test_event_selection_and_roles(client: TestClient, event: int) -> None:
    other = client.post("/api/events", headers=ADMIN, json={"event_code": "VAN26"}).json()["id"]
    assert [e["code"] for e in _report(client)["events"]] == ["LIS26", "VAN26"]
    assert [e["code"] for e in _report(client, events=f"{other},{event}")["events"]] == ["VAN26", "LIS26"]
    client.post(f"/api/documents/{other}/archive", headers=ADMIN)
    assert [e["code"] for e in _report(client)["events"]] == ["LIS26"]  # archived left out by default…
    assert _report(client, events=str(other))["events"][0]["archived"]  # …but can be chosen
    assert any(a["code"] == "VAN26" and a["archived"] for a in _report(client)["available"])
    empty = _report(client, events=str(other))["events"][0]
    assert empty["publishing"]["first_to_last"] is None and empty["requests"]["on_time_pct"] is None

    assert client.get("/api/reports", headers={}).status_code == 403
    assert client.get("/api/reports", headers=VIEWER, params={"events": "999"}).status_code == 404
    assert client.get("/api/reports", headers=VIEWER, params={"events": "1;drop"}).status_code == 422
    tpl = client.post("/api/templates", headers=ADMIN, json={"name": "LIS", "event_id": event}).json()["id"]
    assert client.get("/api/reports", headers=VIEWER, params={"events": str(tpl)}).status_code == 404  # not an event


def test_export(client: TestClient, event: int) -> None:
    client.post("/api/events", headers=ADMIN, json={"event_code": "VAN26"})
    r = client.get("/api/reports/export", headers=VIEWER)
    assert r.status_code == 200 and 'filename="ICT_report.xlsx"' in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]] == ["Measure", "LIS26", "VAN26"]
    labels = [row[0].value for row in ws.iter_rows(min_row=2)]
    for section in ("Cable planning", "ICT request adoption", "Live operations", "Change and publishing"):
        assert section in labels
    points = next(row for row in ws.iter_rows(min_row=2) if row[0].value == "Ethernet points")
    assert [c.value for c in points[1:]] == [126, 0]  # both venues of the fixture


def test_metadata_notes_and_csv(client: TestClient, event: int) -> None:
    fresh = client.post("/api/events", headers=ADMIN, json={"event_code": "VAN26"}).json()["id"]
    report = _report(client, events=f"{event},{fresh}")
    assert report["generated_at"] and "cable_planning" in report["sources"]
    assert "Points planned, not ports used" in report["definitions"]["Ethernet points"]
    van = report["events"][1]
    assert van["created_at"] and any("No ICT request deadline" in n for n in van["notes"])
    assert any("Nothing published" in n for n in van["notes"])

    r = client.get("/api/reports/export", headers=VIEWER, params={"events": f"{event},{fresh}", "format": "csv"})
    assert r.status_code == 200 and 'filename="ICT_report.csv"' in r.headers["content-disposition"]
    text = r.content.decode("utf-8-sig")
    assert text.splitlines()[0] == "Section,Measure,LIS26,VAN26" and "Generated (UTC)" in text
    about = load_workbook(io.BytesIO(client.get("/api/reports/export", headers=VIEWER).content))["About"]
    labels = [row[0].value for row in about.iter_rows()]
    assert "Definitions" in labels and "Sources" in labels
