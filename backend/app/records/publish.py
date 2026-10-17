"""Publish an event's records as the next version: an Excel file in the familiar
workbook layout plus a designed PDF report (Chromium via Gotenberg).

The Excel file re-imports to the same records (records/importer.py), so it can
be shared, edited elsewhere and brought back if needed.
"""

from __future__ import annotations

import hashlib
import html
import io
import logging
import sqlite3
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation

from .. import db
from ..converter import ConversionError, Converter
from ..service import FULL_PDF, ORIGINAL, allocate_next_version
from ..storage import Storage, version_key
from . import changelog, service, store
from .models import LABELS

log = logging.getLogger(__name__)

FONT = "Overpass"
BLACK = PatternFill(fill_type="solid", fgColor="FF1D1C1D")
GREY = PatternFill(fill_type="solid", fgColor="FF666666")
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFFFF")
BODY_FONT = Font(name=FONT, size=10)

WIFI_FIELDS = ["ref_id", "ssid", "coverage", "password_required", "broadcast", "vlan", "band", "password", "purpose"]
WIFI_WIDTHS = [6, 28, 52, 13, 12, 24, 16, 20, 40]
ETHERNET_FIELDS = ["ref_id", "purpose", "location", "vlan", "room", "bandwidth_ns", "bandwidth_ew", "capacity", "ports",
                   "notes"]
ETHERNET_WIDTHS = [6, 30, 30, 28, 14, 18, 18, 12, 12, 58]
PARTNER_FIELDS = ["ref_id", "partner", "pavilion", "stand", "solution", "bandwidth_mbps", "default_cables", "total_ports",
                  "wifi", "vlans", "notes", "purchased"]
PARTNER_WIDTHS = [6, 40, 8, 14, 42, 18, 14, 18, 8, 8, 40, 12]


def _col(fields: list[str], field: str) -> int:
    """1-based spreadsheet column of a field."""
    return fields.index(field) + 1
SKIPPED_ACTIONS = ("reveal", "legacy", "publish", "request", "review")  # not changes to the requirements


# --- snapshot --------------------------------------------------------------------


def snapshot(conn: sqlite3.Connection, doc: int) -> dict[str, Any]:
    """Everything a published version contains, with group/venue names resolved."""
    venues = store.venues(conn, doc)
    groups = store.groups(conn, doc)
    group_name = {g["id"]: g["name"] for g in groups}
    venue_name = {v["id"]: v["name"] for v in venues}
    records = []
    rows = store.records(conn, doc)
    numbers = service.row_numbers(rows)  # IDs are row numbers in display order
    for r in rows:
        data = dict(r["data"]) | ({"ref_id": numbers[r["id"]]} if r["id"] in numbers else {})
        records.append({"id": r["id"], "type": r["type"], "venue": venue_name.get(r["venue_id"]),
                        "group": group_name.get(r["group_id"]), "data": data})
    return {
        "venues": [v["name"] for v in venues],
        "groups": {kind: [g["name"] for g in groups if g["kind"] == kind] for kind in ("wifi", "ethernet")},
        "lists": {kind: [i["value"] for i in store.lists(conn, doc) if i["kind"] == kind]
                  for kind in ("building", "location", "bandwidth", "vlan")},
        "records": records,
        "totals": service.totals(conn, doc),
    }


def _rows(snap: dict, record_type: str, venue: str | None = None, group: str | None = None) -> list[dict]:
    return [r["data"] for r in snap["records"]
            if r["type"] == record_type and (venue is None or r["venue"] == venue)
            and (group is None or r["group"] == group)]


def _display(field: str, value: Any) -> Any:
    if value is None:
        return None
    if field == "purchased":
        return {"yes": "Yes", "no": "No", "tbc": "TBC"}.get(value, value)
    return value


# --- Excel ---------------------------------------------------------------------------


def _header(ws, labels: list[str], widths: list[int]) -> None:
    for i, (label, width) in enumerate(zip(labels, widths), start=1):
        cell = ws.cell(1, i, label)
        cell.font, cell.fill = HEADER_FONT, BLACK
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"


def _band(ws, row: int, name: str, fill: PatternFill, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill, cell.font = fill, HEADER_FONT
    ws.cell(row, 1, name)


def _body(ws, row: int, values: list[Any]) -> None:
    for c, value in enumerate(values, start=1):
        cell = ws.cell(row, c, value)
        cell.font = BODY_FONT


def _dropdown(ws, formula: str, column: int, first: int, last: int) -> None:
    if last < first:
        return
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    letter = get_column_letter(column)
    dv.add(f"{letter}{first}:{letter}{last}")
    ws.add_data_validation(dv)


def build_workbook(snap: dict, log: dict, *, version: int) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    setup_rows = {kind: len(values) + 1 for kind, values in snap["lists"].items()}
    vlan_source = f"{quote_sheetname('Setup')}!$D$2:$D${max(setup_rows['vlan'], 2)}"
    location_source = f"{quote_sheetname('Setup')}!$B$2:$B${max(setup_rows['location'], 2)}"

    for venue in snap["venues"]:
        ws = wb.create_sheet(f"{venue} - Wi-Fi Req"[:31])
        _header(ws, [LABELS["wifi"][f] for f in WIFI_FIELDS], WIFI_WIDTHS)
        row = 2
        for group in snap["groups"]["wifi"]:
            items = _rows(snap, "wifi", venue, group)
            _band(ws, row, group, BLACK, len(WIFI_FIELDS))
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(WIFI_FIELDS))
            row += 1
            for data in items:
                _body(ws, row, [_display(f, data.get(f)) for f in WIFI_FIELDS])
                row += 1
        _dropdown(ws, vlan_source, _col(WIFI_FIELDS, "vlan"), 2, row - 1)

        ws = wb.create_sheet(f"{venue} Ethernet Req"[:31])
        points_col = _col(ETHERNET_FIELDS, "ports")
        letter = get_column_letter(points_col)
        _header(ws, [LABELS["ethernet"][f] for f in ETHERNET_FIELDS], ETHERNET_WIDTHS)
        row, totals = 2, []
        for section in snap["groups"]["ethernet"]:
            items = _rows(snap, "ethernet", venue, section)
            _band(ws, row, section, GREY, len(ETHERNET_FIELDS))
            header_row = row
            row += 1
            for data in items:
                _body(ws, row, [_display(f, data.get(f)) for f in ETHERNET_FIELDS])
                row += 1
            points = ws.cell(header_row, points_col)
            points.value = f"=SUM({letter}{header_row + 1}:{letter}{row - 1})" if items else '=N("section total")'
            points.alignment = Alignment(horizontal="center")
            totals.append(f"{letter}{header_row}")
        _band(ws, row, "Total", GREY, len(ETHERNET_FIELDS))
        ws.cell(row, points_col, f"=SUM({','.join(totals)})" if totals else 0).alignment = Alignment(horizontal="center")
        _dropdown(ws, vlan_source, _col(ETHERNET_FIELDS, "vlan"), 2, row - 1)
        _dropdown(ws, location_source, _col(ETHERNET_FIELDS, "location"), 2, row - 1)

    ws = wb.create_sheet("Partners Custom Orders")
    _header(ws, [LABELS["partner"][f] for f in PARTNER_FIELDS], PARTNER_WIDTHS)
    for i, data in enumerate(_rows(snap, "partner"), start=2):
        _body(ws, i, [_display(f, data.get(f)) for f in PARTNER_FIELDS])
    _dropdown(ws, '"Yes,No,TBC"', _col(PARTNER_FIELDS, "purchased"), 2, max(2, len(_rows(snap, "partner")) + 20))

    ws = wb.create_sheet("Changes - Log")
    _header(ws, ["Change", "Type", "Item", "Where", "Details"], [12, 22, 34, 34, 80])
    rows = changelog.lines(log)
    if log.get("note"):
        rows.insert(0, ("Note", "", "", "", log["note"]))
    if log["first"]:
        rows.insert(1 if log.get("note") else 0, ("First version", "", f"{log['items']} items", "", ""))
    for i, row in enumerate(rows, start=2):
        _body(ws, i, list(row))

    ws = wb.create_sheet("Setup")
    for col, (kind, title) in enumerate([("building", "Building"), ("location", "Hall/Location/Area"),
                                         ("bandwidth", "Bandwidth Allocation"), ("vlan", "VLANs")], start=1):
        head = ws.cell(1, col, title)
        head.font = Font(name=FONT, size=10, bold=True)
        ws.column_dimensions[get_column_letter(col)].width = [32, 45, 20, 24][col - 1]
        for r, value in enumerate(snap["lists"][kind], start=2):
            ws.cell(r, col, value).font = BODY_FONT

    wb.properties.title = f"Event Network Requirements V{version}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --- PDF report -------------------------------------------------------------------------

REPORT_CSS = """
@page { size: A4 landscape; margin: 12mm 10mm 14mm; }
* { box-sizing: border-box; }
body { font-family: 'Overpass', 'Lato', sans-serif; font-size: 8.5pt; color: #111; margin: 0; }
h1 { font-size: 20pt; margin: 0 0 4pt; }
h2 { font-size: 13pt; margin: 0 0 6pt; break-after: avoid; }
h3 { font-size: 10pt; margin: 10pt 0 4pt; break-after: avoid; }
.cover { height: 170mm; display: flex; flex-direction: column; justify-content: center; }
.meta { color: #555; font-size: 10pt; margin: 2pt 0; }
.summary { margin-top: 18pt; border-collapse: collapse; }
.summary td, .summary th { padding: 3pt 10pt 3pt 0; text-align: left; font-size: 9.5pt; }
section { break-before: page; }
table.data { width: 100%; border-collapse: collapse; table-layout: fixed; }
table.data thead { display: table-header-group; }
table.data th { background: #1d1c1d; color: #fff; text-align: left; padding: 4pt 4pt; font-size: 8pt; vertical-align: bottom; }
table.data td { padding: 3pt 4pt; border-bottom: 0.5pt solid #ddd; vertical-align: top; overflow-wrap: break-word; }
table.data th + th { padding-left: 6pt; }
table.data tr { break-inside: avoid; }
tr.group td { background: #1d1c1d; color: #fff; font-weight: 700; }
tr.section td { background: #666; color: #fff; font-weight: 700; }
tr.total td { background: #222; color: #fff; font-weight: 700; }
td.num { text-align: right; }
.muted { color: #777; }
.changes li { margin: 1pt 0; }
"""


def _e(value: Any) -> str:
    return "" if value is None else html.escape(str(value))


def _table(headers: list[str], widths: list[int], body: str, extra_class: str = "") -> str:
    total = sum(widths)
    cols = "".join(f'<col style="width:{w / total * 100:.2f}%">' for w in widths)
    head = "".join(f"<th>{_e(h)}</th>" for h in headers)
    return (f'<table class="data {extra_class}"><colgroup>{cols}</colgroup>'
            f"<thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>")


def _changes_html(log: dict) -> str:
    """The version's changelog section of the report."""
    parts = []
    if log.get("note"):
        parts.append(f'<p class="note"><strong>Note:</strong> {_e(log["note"])}</p>')
    if log["first"]:
        parts.append(f"<p>First version: {log['items']} items.</p>")
    else:
        c = log["counts"]
        parts.append(f"<p>{c['added']} added · {c['removed']} removed · {c['changed']} changed</p>"
                     if any(c.values()) else '<p class="muted">No changes to the requirements since the previous version.</p>')
        for s in log["sections"]:
            items = [f'<li><strong>Added</strong> {_e(i["name"])}<span class="muted"> · {_e(i["where"])}</span>'
                     f'{" — " + _e(i["summary"]) if i["summary"] else ""}</li>' for i in s["added"]]
            items += [f'<li><strong>Removed</strong> {_e(i["name"])}<span class="muted"> · {_e(i["where"])}</span></li>'
                      for i in s["removed"]]
            items += [f'<li><strong>Changed</strong> {_e(i["name"])}<span class="muted"> · {_e(i["where"])}</span>: '
                      + "; ".join(f'{_e(f["label"])} {_e(f["old"])} → {_e(f["new"])}' if f["old"] else
                                  f'{_e(f["label"])}: {_e(f["new"])}' for f in i["fields"]) + "</li>" for i in s["changed"]]
            parts.append(f'<h3>{_e(s["title"])}</h3><ul class="changes">{"".join(items)}</ul>')
    if log["points"] and not log["first"]:
        parts.append("<h3>Ethernet points</h3><ul class=\"changes\">" + "".join(
            f'<li>{_e(p["venue"])}: {_e(p["old"] if p["old"] is not None else "—")} → {_e(p["new"])}</li>'
            for p in log["points"]) + "</ul>")
    return "".join(parts)


def build_report(snap: dict, log: dict, *, title: str, version: int, user: str, when: str) -> str:
    totals = {t["venue"]: t for t in snap["totals"]["ethernet"]}
    count = lambda t, v=None: len(_rows(snap, t, v))  # noqa: E731
    summary = "".join(
        f"<tr><th>{_e(v)}</th><td>{count('wifi', v)} Wi-Fi networks</td>"
        f"<td>{count('ethernet', v)} Ethernet requirements</td><td>{_e(totals[v]['total'])} Ethernet points</td></tr>"
        for v in snap["venues"]
    )
    parts = [f"""<div class="cover"><h1>{_e(title)}</h1>
        <p class="meta">Network requirements · Version {version}</p>
        <p class="meta">Published by {_e(user)} on {_e(when)}</p>
        <table class="summary">{summary}<tr><th>All venues</th><td colspan="3">{count('partner')} partner orders</td></tr></table></div>"""]

    before_points = ETHERNET_FIELDS.index("ports")
    for venue in snap["venues"]:
        rows = []
        for group in snap["groups"]["wifi"]:
            items = _rows(snap, "wifi", venue, group)
            if not items:
                continue
            rows.append(f'<tr class="group"><td colspan="{len(WIFI_FIELDS)}">{_e(group)}</td></tr>')
            rows += ["<tr>" + "".join(f"<td>{_e(_display(f, d.get(f)))}</td>" for f in WIFI_FIELDS) + "</tr>" for d in items]
        body = "".join(rows) or f'<tr><td colspan="{len(WIFI_FIELDS)}" class="muted">No Wi-Fi networks</td></tr>'
        parts.append(f"<section><h2>{_e(venue)} · Wi-Fi networks</h2>"
                     + _table([LABELS["wifi"][f] for f in WIFI_FIELDS], WIFI_WIDTHS, body) + "</section>")

        rows = []
        for section in totals[venue]["sections"]:
            items = _rows(snap, "ethernet", venue, section["section"])
            rows.append(f'<tr class="section"><td colspan="{before_points}">{_e(section["section"])}</td>'
                        f'<td class="num">{_e(section["ports"])}</td><td></td></tr>')
            rows += ["<tr>" + "".join(
                f'<td class="num">{_e(d.get(f))}</td>' if f == "ports" else f"<td>{_e(d.get(f))}</td>"
                for f in ETHERNET_FIELDS) + "</tr>" for d in items]
        rows.append(f'<tr class="total"><td colspan="{before_points}">Total</td><td class="num">{_e(totals[venue]["total"])}</td><td></td></tr>')
        parts.append(f"<section><h2>{_e(venue)} · Ethernet requirements</h2>"
                     + _table([LABELS["ethernet"][f] for f in ETHERNET_FIELDS], ETHERNET_WIDTHS, "".join(rows))
                     + "</section>")

    partner_rows = "".join(
        "<tr>" + "".join(
            f'<td class="num">{_e(_display(f, d.get(f)))}</td>' if f in ("bandwidth_mbps", "default_cables",
                                                                       "total_ports", "wifi", "vlans")
            else f"<td>{_e(_display(f, d.get(f)))}</td>" for f in PARTNER_FIELDS) + "</tr>"
        for d in _rows(snap, "partner")
    ) or f'<tr><td colspan="{len(PARTNER_FIELDS)}" class="muted">No partner orders</td></tr>'
    parts.append("<section><h2>Partner custom orders</h2>"
                 + _table([LABELS["partner"][f] for f in PARTNER_FIELDS], PARTNER_WIDTHS, partner_rows) + "</section>")

    parts.append(f"<section><h2>Changes in this version</h2>{_changes_html(log)}</section>")

    return (f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{_e(title)} V{version}</title>'
            f"<style>{REPORT_CSS}</style></head><body>{''.join(parts)}</body></html>")


# --- publish -------------------------------------------------------------------------------


def changes_since_last_publish(conn: sqlite3.Connection, doc: int) -> list[dict]:
    after = store.last_publish_log_id(conn, doc) or 0
    entries = store.history(conn, doc, after_id=after, limit=10_000)
    return [e for e in reversed(entries) if e["action"] not in SKIPPED_ACTIONS]


def publish(conn: sqlite3.Connection, storage: Storage, converter: Converter, *, doc: int,
            user: str, note: str | None = None) -> tuple[dict[str, Any], str | None]:
    """Create the next version from the records. Returns (version row, failure reason or None)."""
    document = db.get_document(conn, doc)
    if document["kind"] == "template":
        raise service.RecordError(409, "Templates aren't published. Create an event from it and publish that.")
    if not store.has_records(conn, doc):
        raise service.RecordError(409, "There is nothing to publish yet. Import or add requirements first.")
    snap = snapshot(conn, doc)
    if note:
        snap["note"] = note
    previous = store.previous_snapshot(conn, doc)
    changes = changelog.diff(previous[1] if previous else None, snap)
    version_hint = db.max_version(conn, doc) + 1
    xlsx = build_workbook(snap, changes, version=version_hint)
    version = allocate_next_version(conn, document_id=doc, sha256=hashlib.sha256(xlsx).hexdigest(), user=user,
                                    warnings=[], edited_from=None, source="publish")
    n = version["version"]
    if n != version_hint:  # someone published or uploaded at the same moment
        xlsx = build_workbook(snap, changes, version=n)
    when = datetime.now(timezone.utc).strftime("%d %b %Y %H:%M UTC")
    failure = None
    try:
        storage.put(version_key(doc, n, ORIGINAL), xlsx)
        pdf = converter.html_to_pdf(build_report(snap, changes, title=document["name"], version=n, user=user, when=when))
        storage.put(version_key(doc, n, FULL_PDF), pdf)
        db.set_version_status(conn, version["id"], db.STATUS_READY)
    except Exception as exc:
        failure = str(exc) if isinstance(exc, ConversionError) else "Internal error while building the PDF"
        log.exception("publishing document %s V%s failed", doc, n)
        db.set_version_status(conn, version["id"], db.STATUS_FAILED, error=failure)
    with conn:
        store.save_snapshot(conn, version_id=version["id"], doc=doc, data=snap)
        if failure is None:
            store.log(conn, doc=doc, entity="event", action="publish", user=user, new=n,
                      summary=f"Published V{n}: " + ("first version" if changes["first"] else
                              "{added} added, {removed} removed, {changed} changed".format(**changes["counts"]))
                              + (f" — {note}" if note else ""))
    return db.get_version_by_id(conn, version["id"]), failure  # type: ignore[return-value]


def changelog_between(conn: sqlite3.Connection, doc: int, *, to: str, from_version: int | None = None
                      ) -> dict[str, Any]:
    """Changelog for a published version (against the one before it, or ``from_version``), or for
    the unpublished draft (to='draft') against the latest published version."""
    if to == "draft":
        new, to_label = snapshot(conn, doc), "draft"
        previous = store.previous_snapshot(conn, doc)
    else:
        n = int(to)
        new = store.snapshot_for(conn, doc, n)
        if new is None:
            raise service.RecordError(422, f"V{n} wasn't published from the requirements, so it has no item list "
                                           "to compare. Uploaded and edited workbooks can only be viewed.")
        to_label = n
        previous = store.previous_snapshot(conn, doc, before=n)
    if from_version is not None:
        old = store.snapshot_for(conn, doc, from_version)
        if old is None:
            raise service.RecordError(422, f"V{from_version} wasn't published from the requirements, so it can't be compared.")
        previous = (from_version, old)
    log = changelog.diff(previous[1] if previous else None, new)
    return log | {"from_version": previous[0] if previous else None, "to_version": to_label,
                  "published_versions": store.published_versions(conn, doc)}
