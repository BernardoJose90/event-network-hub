from __future__ import annotations

import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import db
from app.config import Settings
from app.main import create_app

from .conftest import FakeConverter
from .records_fixture import records_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ADMIN = {"X-Role": "network-admin", "X-User": "alice"}
VIEWER = {"X-Role": "viewer", "X-User": "victor"}


@pytest.fixture
def client(tmp_path: Path):
    settings = Settings(data_dir=tmp_path, gotenberg_url="http://unused", gotenberg_timeout_seconds=1,
                        max_upload_bytes=1024 * 1024)
    app = create_app(settings)
    with TestClient(app) as c:
        app.state.converter = FakeConverter()
        c.settings = settings
        yield c


@pytest.fixture
def event(client: TestClient) -> int:
    r = client.post("/api/documents/upload", headers=ADMIN,
                    files={"file": ("LIS26 _ Event_Network_Requirements_V1.xlsx", records_workbook(), XLSX)})
    doc = r.json()["document"]["id"]
    assert client.post(f"/api/events/{doc}/import", headers=ADMIN, json={}).status_code == 200
    return doc


def _records(client, doc, rtype):
    return client.get(f"/api/events/{doc}/records/{rtype}", headers=ADMIN).json()


def _dep(client, doc, **data):
    body = {"area": "Main stage", "service": "Live stream encoders", "team": "AV", "criticality": "P1",
            "accountable": "Head of AV", "responsible": "network", "vlans": ["Operational_Prod"],
            "runbook": "Switch the encoders to the 4G bonding unit", "phase": ["live"]} | data
    r = client.post(f"/api/events/{doc}/dependencies", headers=ADMIN, json={"data": body})
    assert r.status_code == 201, r.text
    return r.json()


def test_create_read_and_validate(client: TestClient, event: int) -> None:
    wifi = _records(client, event, "wifi")[0]
    dep = _dep(client, event, records=[wifi["id"]])
    assert dep["data"]["criticality"] == "P1" and dep["linked"] == [{"id": wifi["id"], "type": "wifi",
                                                                    "name": wifi["data"]["ssid"]}]
    assert dep["last_reviewed_by"] == "alice" and not dep["needs_review"] and not dep["missing_owner"]
    listing = client.get(f"/api/events/{event}/dependencies", headers=VIEWER).json()
    assert [d["id"] for d in listing["entries"]] == [dep["id"]] and "AV" in listing["team_suggestions"]

    bad = lambda **d: client.post(f"/api/events/{event}/dependencies", headers=ADMIN,  # noqa: E731
                                  json={"data": {"area": "Hall", "service": "x"} | d})
    assert bad(vlans=["Nope"]).status_code == 422 and "VLAN 'Nope'" in bad(vlans=["Nope"]).json()["detail"]
    assert bad(records=[999999]).status_code == 422
    assert bad(criticality="P9").status_code == 422
    assert bad(area="").status_code == 422
    assert bad(surprise=1).status_code == 422
    assert client.post(f"/api/events/{event}/dependencies", headers=VIEWER,
                       json={"data": {"area": "x", "service": "y"}}).status_code == 403
    no_owner = _dep(client, event, accountable=None, area="Registration", service="Badge printers")
    assert no_owner["missing_owner"]


def test_update_review_delete_and_history(client: TestClient, event: int) -> None:
    dep = _dep(client, event)
    url = f"/api/events/{event}/dependencies/{dep['id']}"
    r = client.put(url, headers=ADMIN, json={"revision": dep["revision"], "data": dep["data"] | {"criticality": "P2"}})
    assert r.status_code == 200 and r.json()["data"]["criticality"] == "P2"
    assert client.put(url, headers=ADMIN, json={"revision": dep["revision"], "data": dep["data"]}).status_code == 409
    dep = r.json()

    conn = db.connect(client.settings.db_path)
    with conn:
        conn.execute("UPDATE event_dependencies SET last_reviewed_at = '2020-01-01T00:00:00+00:00' WHERE id = ?",
                     (dep["id"],))
    conn.close()
    stale = client.get(f"/api/events/{event}/dependencies", headers=ADMIN).json()["entries"][0]
    assert stale["needs_review"]
    r = client.post(f"{url}/review", headers=ADMIN, json={"revision": stale["revision"]})
    assert r.status_code == 200 and not r.json()["needs_review"]

    assert client.delete(f"{url}?revision={r.json()['revision']}", headers=ADMIN).status_code == 204
    log = client.get(f"/api/events/{event}/history?entity=dependency", headers=ADMIN).json()
    assert [e["action"] for e in log] == ["delete", "review", "update", "create"]
    assert log[2]["field"] == "criticality" and (log[2]["old"], log[2]["new"]) == ("P1", "P2")


def test_vlan_rename_follows_and_delete_is_blocked(client: TestClient, event: int) -> None:
    dep = _dep(client, event)
    lists = client.get(f"/api/events/{event}/overview", headers=ADMIN).json()["lists"]
    prod = next(i for i in lists if i["value"] == "Operational_Prod")
    client.patch(f"/api/events/{event}/lists/item/{prod['id']}", headers=ADMIN, json={"value": "Ops_Prod"})
    entry = client.get(f"/api/events/{event}/dependencies", headers=ADMIN).json()["entries"][0]
    assert entry["data"]["vlans"] == ["Ops_Prod"]
    # still used by records too; remove those first, then the register blocks it
    for r in _records(client, event, "wifi") + _records(client, event, "ethernet"):
        if r["data"].get("vlan") == "Ops_Prod":
            client.patch(f"/api/events/{event}/records/{r['type']}/{r['id']}", headers=ADMIN,
                         json={"revision": r["revision"], "data": {"vlan": None}})
    r = client.delete(f"/api/events/{event}/lists/item/{prod['id']}", headers=ADMIN)
    assert r.status_code == 409 and "dependency register" in r.json()["detail"]
    assert dep["id"] == entry["id"]


def test_copied_into_templates_and_new_events(client: TestClient, event: int) -> None:
    wifi = _records(client, event, "wifi")
    groups = {g["id"]: g["name"] for g in client.get(f"/api/events/{event}/overview", headers=ADMIN).json()["groups"]}
    partner_net = next(w for w in wifi if groups[w["group_id"]] == "Partners")
    core = next(w for w in wifi if groups[w["group_id"]] != "Partners")
    _dep(client, event, records=[core["id"], partner_net["id"]])
    tpl = client.post("/api/templates", headers=ADMIN, json={"name": "LIS", "event_id": event}).json()["id"]
    entry = client.get(f"/api/events/{tpl}/dependencies", headers=ADMIN).json()["entries"][0]
    assert entry["data"]["service"] == "Live stream encoders"
    assert [l["name"] for l in entry["linked"]] == [core["data"]["ssid"]]  # partner network not in templates
    assert entry["linked"][0]["id"] != core["id"]  # points at the template's own copy

    new = client.post("/api/events", headers=ADMIN, json={"event_code": "LIS27", "template_id": tpl}).json()["id"]
    copied = client.get(f"/api/events/{new}/dependencies", headers=ADMIN).json()["entries"]
    assert len(copied) == 1 and copied[0]["linked"][0]["name"] == core["data"]["ssid"]


def test_published_with_the_requirements(client: TestClient, event: int) -> None:
    _dep(client, event, runbook="Use <the> 4G unit")
    assert client.post(f"/api/events/{event}/publish", headers=ADMIN).status_code == 201
    html = client.app.state.converter.html
    assert "Network dependencies" in html and "Live stream encoders" in html and "&lt;the&gt;" in html
    from app.storage import version_key

    xlsx = client.app.state.storage.get(version_key(event, 2, "original.xlsx"))
    ws = load_workbook(io.BytesIO(xlsx))["Network dependencies"]
    assert [c.value for c in ws[2]][:4] == ["Main stage", "Live stream encoders", "AV", "P1"]


def test_reimport_keeps_register_but_drops_links(client: TestClient, event: int) -> None:
    wifi = _records(client, event, "wifi")[0]
    _dep(client, event, records=[wifi["id"]])
    r = client.post(f"/api/events/{event}/import", headers=ADMIN, json={"replace": True})
    assert r.status_code == 200
    entry = client.get(f"/api/events/{event}/dependencies", headers=ADMIN).json()["entries"][0]
    assert entry["data"]["records"] == [] and entry["data"]["service"] == "Live stream encoders"


def test_archived_events_are_read_only(client: TestClient, event: int) -> None:
    dep = _dep(client, event)
    client.post(f"/api/documents/{event}/archive", headers=ADMIN)
    assert client.get(f"/api/events/{event}/dependencies", headers=VIEWER).status_code == 200
    r = client.put(f"/api/events/{event}/dependencies/{dep['id']}", headers=ADMIN,
                   json={"revision": dep["revision"], "data": dep["data"]})
    assert r.status_code == 409


def test_help_desk_ticket_shows_dependencies(client: TestClient, event: int) -> None:
    _dep(client, event, area="Main stage", criticality="P1")
    _dep(client, event, area="Main Stage backstage", service="Talkback", criticality="P2")
    _dep(client, event, area="Hall 9", service="CCTV")
    t = client.post("/api/tickets", headers=ADMIN, json={"title": "No network", "category": "wired",
                                                         "event_id": event, "location": "main  stage"}).json()
    r = client.get(f"/api/tickets/{t['id']}/dependencies", headers=VIEWER).json()
    assert sorted(d["data"]["service"] for d in r["entries"]) == ["Live stream encoders", "Talkback"]
    assert r["suggested_severity"] == "P1"
    other = client.post("/api/tickets", headers=ADMIN, json={"title": "x", "category": "wifi"}).json()
    assert client.get(f"/api/tickets/{other['id']}/dependencies", headers=VIEWER).json() == {
        "entries": [], "suggested_severity": None}
