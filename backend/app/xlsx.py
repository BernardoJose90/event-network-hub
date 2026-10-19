"""Workbook inspection with openpyxl."""

from __future__ import annotations

import io

from openpyxl import load_workbook


def visible_sheet_count(data: bytes) -> int:
    wb = load_workbook(io.BytesIO(data), read_only=True)
    try:
        return sum(1 for name in wb.sheetnames if wb[name].sheet_state == "visible")
    finally:
        wb.close()
