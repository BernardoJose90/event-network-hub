from __future__ import annotations

from openpyxl import Workbook

from app.xlsx import find_formula_errors

from .conftest import xlsx_bytes


def test_finds_each_error_with_sheet_and_cell() -> None:
    def build(wb: Workbook) -> None:
        ws = wb.active
        ws.title = "Summary"
        ws["A1"] = "ok"
        ws["B2"] = "#REF!"
        ws["C3"] = "#NAME?"
        ws["D4"] = "#DIV/0!"
        other = wb.create_sheet("Detail")
        other["A1"] = "#VALUE!"
        other["E10"] = "#N/A"

    assert find_formula_errors(xlsx_bytes(build)) == [
        {"sheet": "Summary", "cell": "B2", "error": "#REF!"},
        {"sheet": "Summary", "cell": "C3", "error": "#NAME?"},
        {"sheet": "Summary", "cell": "D4", "error": "#DIV/0!"},
        {"sheet": "Detail", "cell": "A1", "error": "#VALUE!"},
        {"sheet": "Detail", "cell": "E10", "error": "#N/A"},
    ]


def test_clean_workbook_has_no_warnings() -> None:
    def build(wb: Workbook) -> None:
        ws = wb.active
        ws.append(["Name", "Total"])
        ws.append(["a", 1])
        ws["C1"] = "=SUM(B2:B2)"

    assert find_formula_errors(xlsx_bytes(build)) == []


def test_ignores_text_that_merely_contains_error_token() -> None:
    def build(wb: Workbook) -> None:
        wb.active["A1"] = "see #REF! in the old file"

    assert find_formula_errors(xlsx_bytes(build)) == []
