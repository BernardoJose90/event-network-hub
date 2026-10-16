from __future__ import annotations

import io
import zipfile

from openpyxl import Workbook, load_workbook

from app.pdf_prep import _strip_print_gridlines, prepare_for_pdf

from .conftest import xlsx_bytes


def _with_print_gridlines(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Setup"
    ws["A1"] = "Building"
    ws.print_options.gridLines = True
    ws.print_options.horizontalCentered = True
    other = wb.create_sheet("Plain")
    other["A1"] = "untouched"


def _sheet_xml(data: bytes, name: str) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return zf.read(name)


def test_print_gridlines_are_removed() -> None:
    out = prepare_for_pdf(xlsx_bytes(_with_print_gridlines))

    ws = load_workbook(io.BytesIO(out))["Setup"]
    assert not ws.print_options.gridLines
    assert ws.print_options.horizontalCentered  # other print options kept
    assert ws["A1"].value == "Building"


def test_other_parts_are_byte_identical() -> None:
    data = xlsx_bytes(_with_print_gridlines)
    out = prepare_for_pdf(data)

    with zipfile.ZipFile(io.BytesIO(data)) as a, zipfile.ZipFile(io.BytesIO(out)) as b:
        assert a.namelist() == b.namelist()
        changed = [n for n in a.namelist() if a.read(n) != b.read(n)]
    assert changed == ["xl/worksheets/sheet1.xml"]


def test_workbook_without_print_gridlines_is_returned_unchanged() -> None:
    data = xlsx_bytes()
    assert prepare_for_pdf(data) is data


def test_strip_handles_attribute_variants() -> None:
    xml = b'<printOptions gridLines="1" gridLinesSet=\'1\' horizontalCentered="1"/>'
    assert _strip_print_gridlines(xml) == b'<printOptions horizontalCentered="1"/>'
    assert _strip_print_gridlines(b'<printOptions/>') == b'<printOptions/>'
