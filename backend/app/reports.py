"""Cross-event reporting for the ICT team: /api/reports.

One set of measures per event, grouped by objective, so events can be compared
side by side ("event over event"):
1. cable planning        — Ethernet points, how many came through ICT requests, what arrived late;
2. ICT request adoption  — requests per team, on time, "not sure" answers, review time;
3. live operations       — the help desk KPIs (see tickets/report.py);
4. change and publishing — versions, and what changed from the first to the last version.
"""

from __future__ import annotations

import csv
import io
import logging
from collections.abc import Iterator
from datetime import datetime
from statistics import mean
from typing import Annotated, Any

from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response, status
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db
from .auth import Principal, current_principal, require_known_role
from .ictrequests import service as requests_service
from .ictrequests.models import REQUEST_TEAMS
from .records import changelog, service as records_service, store as records_store
from .tickets import report as ticket_report, store as ticket_store

log = logging.getLogger(__name__)
router = APIRouter(prefix="/api/reports")

GOALS = {"ack_within_target_pct": 90, "closure_quality_pct": 90, "requests_on_time_pct": 90}

# Where each section's numbers come from, and how each measure is calculated. Shown on the
# page ("Sources and definitions") and in the export, so figures can be read without the code.
SOURCES = {
    "cable_planning": "The event's current Ethernet requirements (including unpublished changes), its approved ICT "
                      "requests, the change history and help desk tickets.",
    "requests": "ICT requests submitted through the team links for the event.",
    "operations": "Help desk tickets filed under the event. Duplicates merged into another ticket are left out.",
    "publishing": "The event's versions and the snapshots saved when each version was published.",
}
DEFINITIONS = {
    "Ethernet points": "Sum of Ethernet points across all venues and sections. Points planned, not ports used: "
                       "ports actually used after the event aren't recorded.",
    "Points via ICT requests": "Points of Ethernet items created by approving an ICT request, and their share of all points.",
    "Added after the deadline": "Ethernet items created after the event's ICT request deadline (by the date they were "
                                "added), and their points. Shown only when a deadline is set.",
    "Wired tickets on site": "Help desk tickets in the Wired connectivity category.",
    "Submitted on time": "ICT requests submitted on or before the deadline, as a share of all requests. Goal 90%.",
    "With 'not sure' answers": "Requests where the team didn't know the number of cables, PoE or speed (wired), or the "
                               "number of devices (Wi-Fi).",
    "Mean review time": "Average time from submission to the first review (approved, declined or asked for more information).",
    "Acknowledged within 5 min": "Tickets acknowledged within 5 minutes of being raised, as a share of tickets whose "
                                 "acknowledge time has passed. Goal 90%.",
    "Closed with root cause and fix": "Resolved or closed tickets with both a root cause and a resolution. Goal 90%.",
    "Recurring issues": "Tickets in the same category and area (stand or location) as an earlier ticket, in any event, "
                        "as a share of tickets with an area.",
    "First → last version": "Item-level changelog between the first and the last version published from the requirements.",
}


def _conn(request: Request) -> Iterator[db.Connection]:
    conn = db.connect(request.app.state.settings.database_url)
    try:
        yield conn
    finally:
        conn.close()


Conn = Annotated[db.Connection, Depends(_conn)]


def _reader(principal: Annotated[Principal, Depends(current_principal)]) -> Principal:
    require_known_role(principal)
    return principal


Reader = Annotated[Principal, Depends(_reader)]


def _pct(part: float, whole: float) -> float | None:
    return round(100 * part / whole, 1) if whole else None


# --- measures -------------------------------------------------------------------------------


def cable_planning(conn: db.Connection, doc: int) -> dict[str, Any]:
    totals = records_service.totals(conn, doc)["ethernet"]
    ethernet = {r["id"]: r for r in records_store.records(conn, doc, "ethernet")}
    sections: dict[str, int | float] = {}
    for venue in totals:
        for s in venue["sections"]:
            sections[s["section"]] = sections.get(s["section"], 0) + s["ports"]
    total = sum(v["total"] for v in totals)
    approved = [r for r in requests_service.list_for(conn, doc) if r["status"] == "approved" and r["kind"] == "wired"]
    via_requests = sum((ethernet[r["record_id"]]["data"].get("ports") or 0) for r in approved if r["record_id"] in ethernet)
    deadline = requests_service.settings(conn, doc)["deadline"]
    late_items, late_points = 0, 0
    if deadline:
        created = {e["record_id"]: e["at"] for e in records_store.history(conn, doc, entity="ethernet", limit=100_000)
                   if e["action"] == "create"}
        for rid, item in ethernet.items():
            if rid in created and created[rid][:10] > deadline:
                late_items += 1
                late_points += item["data"].get("ports") or 0
    wired_tickets = sum(1 for t in ticket_store.list_tickets(conn, event_id=doc, limit=100_000)
                        if t["category"] == "wired" and not t["merged_into"])
    return {
        "points": total,
        "items": len(ethernet),
        "by_venue": [{"venue": v["venue"], "points": v["total"]} for v in totals],
        "by_section": [{"section": k, "points": v} for k, v in sorted(sections.items(), key=lambda x: -x[1]) if v],
        "points_via_requests": via_requests,
        "via_requests_pct": _pct(via_requests, total),
        "deadline": deadline,
        "items_added_after_deadline": late_items if deadline else None,
        "points_added_after_deadline": late_points if deadline else None,
        "wired_tickets_on_site": wired_tickets,
    }


def request_adoption(conn: db.Connection, doc: int) -> dict[str, Any]:
    reqs = requests_service.list_for(conn, doc)
    summary = requests_service.summary(conn, doc)

    def review_hours(items):
        hours = [(datetime.fromisoformat(r["reviewed_at"]) - datetime.fromisoformat(r["submitted_at"])).total_seconds()
                 / 3600 for r in items if r["reviewed_at"]]
        return round(mean(hours), 1) if hours else None

    teams = {}
    for key, label in REQUEST_TEAMS.items():
        mine = [r for r in reqs if r["team"] == key]
        teams[key] = {
            "label": label,
            "requests": len(mine),
            "on_time_pct": _pct(sum(not r["late"] for r in mine), len(mine)),
            "approved": sum(r["status"] == "approved" for r in mine),
            "declined": sum(r["status"] == "declined" for r in mine),
            "open": sum(r["status"] in ("submitted", "needs_info") for r in mine),
            "ports_requested": summary["by_team"][key]["ports_requested"],
            "review_hours": review_hours(mine),
        }
    return {
        "requests": len(reqs),
        "teams_using": sum(1 for t in teams.values() if t["requests"]),
        "on_time_pct": _pct(sum(not r["late"] for r in reqs), len(reqs)),
        "not_sure_pct": _pct(summary["unknowns"], len(reqs)),
        "review_hours": review_hours(reqs),
        "open": summary["by_status"]["submitted"] + summary["by_status"]["needs_info"],
        "by_team": teams,
    }


def live_operations(conn: db.Connection, doc: int) -> dict[str, Any]:
    k = ticket_report.kpis(conn, doc)
    return {
        "tickets": k["tickets"],
        "open": k["open"],
        "ack_within_target_pct": k["acknowledge"]["within_target_pct"],
        "ack_median_minutes": k["acknowledge"]["median_minutes"],
        "closure_quality_pct": k["closure_quality"]["pct"],
        "resolve_mean_minutes": k["resolve"]["mean_minutes"],
        "partner_impacting": k["partner_impacting"]["count"],
        "partner_mean_minutes": k["partner_impacting"]["mean_minutes_to_resolve"],
        "recurrence_pct": k["recurrence"]["pct"],
        "by_category": k["by_category"],
    }


def change_and_publishing(conn: db.Connection, doc: int) -> dict[str, Any]:
    published = records_store.published_versions(conn, doc)
    versions = db.list_versions(conn, doc)
    out: dict[str, Any] = {
        "versions": len(versions),
        "published": len(published),
        "first_published": None, "last_published": None,
        "unpublished_changes": None,
        "first_to_last": None,
    }
    when = {v["version"]: v["uploaded_at"] for v in versions}
    if published:
        out["first_published"] = when.get(published[0])
        out["last_published"] = when.get(published[-1])
    if len(published) >= 2:
        diff = changelog.diff(records_store.snapshot_for(conn, doc, published[0]),
                              records_store.snapshot_for(conn, doc, published[-1]))
        points = sum((p["new"] or 0) - (p["old"] or 0) for p in diff["points"])
        out["first_to_last"] = diff["counts"] | {"from": published[0], "to": published[-1], "points_delta": points}
    last = records_store.last_publish_log_id(conn, doc)
    from .records.publish import SKIPPED_ACTIONS

    out["unpublished_changes"] = sum(1 for e in records_store.history(conn, doc, after_id=last or 0, limit=100_000)
                                     if e["action"] not in SKIPPED_ACTIONS)
    return out


def data_notes(report: dict[str, Any]) -> list[str]:
    """Things to check before comparing this event with others (data quality, not performance)."""
    notes = []
    cp, rq, ops, pub = (report[k] for k in ("cable_planning", "requests", "operations", "publishing"))
    if cp["items"] and not cp["points"]:
        notes.append(f"{cp['items']} Ethernet items but 0 points: points may not have been entered yet "
                     "(events started from a template begin at 0).")
    if not cp["deadline"]:
        notes.append("No ICT request deadline is set, so late additions can't be measured.")
    if not rq["requests"]:
        notes.append("No ICT requests yet.")
    if not ops["tickets"]:
        notes.append("No help desk tickets filed under this event.")
    if not pub["published"]:
        notes.append("Nothing published from the requirements yet.")
    elif pub["unpublished_changes"]:
        notes.append(f"{pub['unpublished_changes']} change(s) not yet published: figures use the current draft.")
    return notes


def event_report(conn: db.Connection, document: dict[str, Any]) -> dict[str, Any]:
    doc = document["id"]
    report = {
        "id": doc,
        "code": document["name"].split(" | ")[0],
        "archived": bool(document["archived_at"]),
        "created_at": document["created_at"],
        "cable_planning": cable_planning(conn, doc),
        "requests": request_adoption(conn, doc),
        "operations": live_operations(conn, doc),
        "publishing": change_and_publishing(conn, doc),
    }
    return report | {"notes": data_notes(report)}


def _events(conn: db.Connection, ids: str | None) -> list[dict[str, Any]]:
    events = db.list_documents(conn)  # events only (templates excluded)
    if not ids:
        return [e for e in events if not e["archived_at"]]
    wanted = []
    for part in ids.split(","):
        if not part.strip().isdigit():
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, "events must be a comma-separated list of ids")
        wanted.append(int(part))
    by_id = {e["id"]: e for e in events}
    missing = [i for i in wanted if i not in by_id]
    if missing:
        raise HTTPException(status.HTTP_404_NOT_FOUND, f"Event {missing[0]} not found")
    return [by_id[i] for i in dict.fromkeys(wanted)]


Events = Annotated[str | None, Query(max_length=500, description="Comma-separated event ids; default: all active")]


@router.get("")
def reports(conn: Conn, principal: Reader, events: Events = None) -> dict[str, Any]:
    chosen = _events(conn, events)
    return {
        "generated_at": db.utcnow(),
        "sources": SOURCES,
        "definitions": DEFINITIONS,
        "events": [event_report(conn, d) for d in chosen],
        "available": [{"id": d["id"], "code": d["name"].split(" | ")[0], "archived": bool(d["archived_at"])}
                      for d in db.list_documents(conn)],
        "goals": GOALS,
        "teams": REQUEST_TEAMS,
    }


# --- export ---------------------------------------------------------------------------------

def _rows(reports_: list[dict[str, Any]]) -> list[tuple[str, list[tuple[str, Any]]]]:
    """(section title, [(measure, getter)]) used for both the sheet rows and their order."""
    def pct(v):
        return None if v is None else f"{v}%"
    return [
        ("Cable planning", [
            ("Ethernet points", lambda r: r["cable_planning"]["points"]),
            ("Ethernet items", lambda r: r["cable_planning"]["items"]),
            ("Points via ICT requests", lambda r: r["cable_planning"]["points_via_requests"]),
            ("Share via ICT requests", lambda r: pct(r["cable_planning"]["via_requests_pct"])),
            ("Request deadline", lambda r: r["cable_planning"]["deadline"]),
            ("Items added after the deadline", lambda r: r["cable_planning"]["items_added_after_deadline"]),
            ("Points added after the deadline", lambda r: r["cable_planning"]["points_added_after_deadline"]),
            ("Wired tickets on site", lambda r: r["cable_planning"]["wired_tickets_on_site"]),
        ]),
        ("ICT request adoption", [
            ("Requests", lambda r: r["requests"]["requests"]),
            ("Teams using the form", lambda r: r["requests"]["teams_using"]),
            ("Submitted on time", lambda r: pct(r["requests"]["on_time_pct"])),
            ("With 'not sure' answers", lambda r: pct(r["requests"]["not_sure_pct"])),
            ("Mean review time (hours)", lambda r: r["requests"]["review_hours"]),
            *[(f"{label}: requests", (lambda k: lambda r: r["requests"]["by_team"][k]["requests"])(key))
              for key, label in REQUEST_TEAMS.items()],
        ]),
        ("Live operations", [
            ("Tickets", lambda r: r["operations"]["tickets"]),
            ("Acknowledged within 5 min", lambda r: pct(r["operations"]["ack_within_target_pct"])),
            ("Median time to acknowledge (min)", lambda r: r["operations"]["ack_median_minutes"]),
            ("Closed with root cause and fix", lambda r: pct(r["operations"]["closure_quality_pct"])),
            ("Mean time to resolve (min)", lambda r: r["operations"]["resolve_mean_minutes"]),
            ("Partner-impacting tickets", lambda r: r["operations"]["partner_impacting"]),
            ("Partner mean time to resolve (min)", lambda r: r["operations"]["partner_mean_minutes"]),
            ("Recurring issues", lambda r: pct(r["operations"]["recurrence_pct"])),
        ]),
        ("Change and publishing", [
            ("Versions", lambda r: r["publishing"]["versions"]),
            ("Published from the requirements", lambda r: r["publishing"]["published"]),
            ("First published", lambda r: (r["publishing"]["first_published"] or "")[:10] or None),
            ("Last published", lambda r: (r["publishing"]["last_published"] or "")[:10] or None),
            ("First → last: added", lambda r: (r["publishing"]["first_to_last"] or {}).get("added")),
            ("First → last: removed", lambda r: (r["publishing"]["first_to_last"] or {}).get("removed")),
            ("First → last: changed", lambda r: (r["publishing"]["first_to_last"] or {}).get("changed")),
            ("First → last: Ethernet points", lambda r: (r["publishing"]["first_to_last"] or {}).get("points_delta")),
            ("Unpublished changes", lambda r: r["publishing"]["unpublished_changes"]),
        ]),
    ]


@router.get("/export")
def export(conn: Conn, principal: Reader, events: Events = None,
           format: Annotated[str, Query(pattern="^(csv|xlsx)$")] = "xlsx") -> Response:
    chosen = [event_report(conn, d) for d in _events(conn, events)]
    generated = db.utcnow()
    if format == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(["Section", "Measure", *[r["code"] for r in chosen]])
        for title, measures in _rows(chosen):
            for label, get in measures:
                writer.writerow([title, label, *[get(r) for r in chosen]])
        writer.writerow([])
        writer.writerow(["Generated (UTC)", generated[:16].replace("T", " ")])
        log.info("%s exported the ICT report as CSV", principal.user)
        return Response(buffer.getvalue().encode("utf-8-sig"), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": 'attachment; filename="ICT_report.csv"'})
    wb = Workbook()
    ws = wb.active
    ws.title = "ICT report"
    navy = PatternFill("solid", fgColor="0C1439")
    teal = PatternFill("solid", fgColor="008484")
    ws.cell(1, 1, "Measure").font = Font(bold=True, color="FFFFFF")
    ws.cell(1, 1).fill = navy
    ws.column_dimensions["A"].width = 38
    for c, r in enumerate(chosen, start=2):
        cell = ws.cell(1, c, r["code"] + (" (archived)" if r["archived"] else ""))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = navy
        ws.column_dimensions[get_column_letter(c)].width = 16
    row = 2
    for title, measures in _rows(chosen):
        cell = ws.cell(row, 1, title)
        cell.font = Font(bold=True, color="FFFFFF")
        for c in range(1, len(chosen) + 2):
            ws.cell(row, c).fill = teal
        row += 1
        for label, get in measures:
            ws.cell(row, 1, label)
            for c, r in enumerate(chosen, start=2):
                ws.cell(row, c, get(r)).alignment = Alignment(horizontal="right")
            row += 1
    ws.freeze_panes = "B2"

    about = wb.create_sheet("About")
    about.column_dimensions["A"].width = 32
    about.column_dimensions["B"].width = 110
    lines = [("Report", "ICT report: events compared against the ICT objectives"),
             ("Generated (UTC)", generated[:16].replace("T", " ")),
             ("Events", ", ".join(r["code"] for r in chosen)), ("", ""), ("Sources", "")]
    lines += [(k.replace("_", " ").capitalize(), v) for k, v in SOURCES.items()]
    lines += [("", ""), ("Definitions", "")] + list(DEFINITIONS.items())
    notes = [(r["code"], n) for r in chosen for n in r["notes"]]
    if notes:
        lines += [("", ""), ("Data notes", "")] + notes
    for i, (k, v) in enumerate(lines, start=1):
        about.cell(i, 1, k).font = Font(bold=True)
        about.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    out = io.BytesIO()
    wb.save(out)
    log.info("%s exported the ICT report for %s", principal.user, ",".join(r["code"] for r in chosen))
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="ICT_report.xlsx"'})
