"""A small Event Network Requirements workbook with the real file's quirks."""

from __future__ import annotations

import io
from functools import lru_cache

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import Color

BLACK = PatternFill(fill_type="solid", fgColor="FF1D1C1D")
THEME_BLACK = PatternFill(fill_type="solid", fgColor=Color(theme=1))  # how Google exports some black rows
GREY = PatternFill(fill_type="solid", fgColor="FF666666")
LIGHT = PatternFill(fill_type="solid", fgColor="FFF3F3F3")
WHITE_BOLD = Font(bold=True, color="FFFFFFFF")


def _section(ws, row: int, name: str, fill: PatternFill) -> None:
    ws.cell(row, 1, name).fill = fill
    ws.cell(row, 1).font = WHITE_BOLD


@lru_cache(maxsize=None)
def records_workbook() -> bytes:
    wb = Workbook()
    wifi_a = wb.active
    wifi_a.title = "Altice Arena - Wi-Fi Req"
    eth_a = wb.create_sheet("Altice Arena - Ethernet Req")
    wifi_f = wb.create_sheet("FIL 1-5 - Wi-Fi Req")
    eth_f = wb.create_sheet("FIL 1-5 Ethernet Req")
    partners = wb.create_sheet("Partners Custom Orders")
    vlans = wb.create_sheet("Extending Vlans")
    changes = wb.create_sheet("Changes - Log")
    setup = wb.create_sheet("Setup")
    wb.create_sheet("Scratch notes")["A1"] = "not part of the requirements"

    wifi_headers = ["ID", "SSID", "Coverage area", "Password?", "Broadcast?", "VLAN", "Band", "Password value", "Purpose"]
    for ws in (wifi_a, wifi_f):
        ws.append(wifi_headers)
    _section(wifi_a, 2, "Core Networks", BLACK)
    wifi_a.append([1, "#WebSummit", "all areas", "open network", "yes", "General", "5 GHz & 6 GHz", "none", "General Access"])
    wifi_a.append([2, "#Production", "offices", "Yes", "yes", "Operational_Prod", "5 GHz", "StaffPass1!", "Staff Network"])
    _section(wifi_a, 5, "Lounges", THEME_BLACK)
    wifi_a.append([9])  # placeholder ID only
    _section(wifi_a, 7, "Partners", THEME_BLACK)
    wifi_a["A8"] = "None at this time – typically < 10 private SSIDs"
    wifi_a["A8"].fill = LIGHT
    _section(wifi_f, 2, "Core Networks", BLACK)
    wifi_f.append([1, "#WebSummit", "all areas", "open network", "yes", "General", "5 GHz", "none", "General Access"])
    _section(wifi_f, 4, "Partners", THEME_BLACK)
    wifi_f.append([13, "Enterprise Georgia", "P3", "yes", "maybe", "Partner_Custom", "5 GHz", "Georgia25", None])

    eth_headers = [143, "Location", "VLAN", "Area/Room Number", "Bandwidth per device\nNorth - South",
                   "Bandwidth per deviceEast \n- West", "Capacity", "Ethernet Points", "Notes"]
    for ws in (eth_a, eth_f):
        ws.append(eth_headers)
    rows_a = [
        ("INTERNAL WORKSPACES", None),
        (["Photographers Office", None, "Operational_Prod", None, "50 Mbps", "1 Gbps", None, 2, None], None),
        (["PCT Event Control", "Tejo Roof Container 1", "Operational_Prod", None, "50 Mbps", "1 Gbps", None, 11,
          "1 drop to CCTV"], None),
        (["FoodTrucks / Bar", None, "Operational_Catering", None, "30 Mbps", "N/A", None, 98, None], None),
        (["Bag Drop 1", None, None, None, None, None, None, None, None], None),
        ("PARTNER MASTERCLASSES", None),
        ("Stages", None),
        (["Stage 1", "Mystery Hall", "Operations_Stages", None, "50 Mbps", "1 Gbps", None, "13", None], None),
        (["Stage 2", None, "Operations_Stages", None, None, None, None, "tbc", None], None),
        ("Total", None),
    ]
    r = 2
    for value, _ in rows_a:
        if isinstance(value, str):
            _section(eth_a, r, value, GREY)
        else:
            for c, v in enumerate(value, start=1):
                eth_a.cell(r, c, v)
        r += 1
    _section(eth_f, 2, "INTERNAL WORKSPACES", GREY)
    eth_f.append(["Site Office", "PAV 1", "Operational_Prod", None, "50 Mbps", "N/A", None, 2, None])

    partners.append(["Partner & Exhibitor", "Pav", "Stand Number", "Network Solution", "Bandwidth Amount",
                     "Default Cable", "Total ethernet ports", "WiFI", "Vlan", "Note(s)", "Purchased"])
    partners.append(["IBM", 1, "E178", "Private Wi-Fi + Switch", 60, 4, None, 1, 1, None, "Yes"])
    partners.append(["Pitchbook", 5, None, "4 Ethernet drops", 10, 4, None, None, None, "investor area", None])
    partners.append(["Visa, Inc", 1, "E179", "Private-WiFi+50Mbps", "fifty", 0, None, 1, 1, None, "TBC"])

    vlans.append(["Vlan name", "FIL", "Altice Arena"])
    vlans.append(["Operational_Prod", True, True])
    vlans.append(["General", False, True])
    vlans.append([None, False, False])

    changes.append(["Area", "Change Qty", "V0", "V1", "Reason (latestes)"])
    changes.append(["Executive lounge", "=D2-C2", 0, 2, "Added 2 turnstiles"])
    changes.append(["Password for #Stage changed", 0, 1, 1, "Changed"])

    setup["A1"], setup["B1"], setup["C1"], setup["D1"] = "Building", "Hall/Location/Area", "Bandwidth Allocation", "VLANs"
    for cell in setup[1]:
        cell.font = Font(bold=True)
    buildings = ["MEO Arena", "PAV 1", "Tejo Roof Container 1"]
    locs = ["MEO Arena", "PAV 1", "Tejo Roof Container 1", "PT Meeting Space"]
    bandwidths = ["10Mbs", "20Mbs"]
    vlan_values = ["General", "Operational_Prod", "Operations_Stages", "Partner_Custom"]  # Catering missing
    for i in range(max(len(buildings), len(locs), len(vlan_values))):
        row = [buildings[i] if i < len(buildings) else None, locs[i] if i < len(locs) else None,
               bandwidths[i] if i < len(bandwidths) else None, vlan_values[i] if i < len(vlan_values) else None]
        setup.append(row)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
