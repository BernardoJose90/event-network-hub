from __future__ import annotations

import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import db
from app.config import Settings
from app.main import create_app

from .conftest import FakeConverter
from .sheet_fixture import sheet_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ADMIN = {"X-Role": "network-admin", "X-User": "alice"}


@pytest.fixture
def client(tmp_path: Path):
    settings = Settings(data_dir=tmp_path, gotenberg_url="http://unused", gotenberg_timeout_seconds=1,
                        max_upload_bytes=1024 * 1024)
    app = create_app(settings)
    with TestClient(app) as c:
        app.state.converter = FakeConverter()
        c.settings = settings
        yield c


def _upload(client: TestClient, version: int = 1) -> int:
    r = client.post(
        "/api/documents/upload",
        files={"file": (f"LIS26 _ Event_Network_Requirements_V{version}.xlsx", sheet_workbook(), XLSX)},
        headers=ADMIN,
    )
    assert r.status_code == 201, r.text
    return r.json()["document"]["id"]


def _start(client: TestClient, doc_id: int) -> dict:
    r = client.post(f"/api/documents/{doc_id}/sheet", headers=ADMIN)
    assert r.status_code == 200, r.text
    return r.json()


def _edit_first_cell(data: dict, text: str) -> dict:
    eth = next(s for s in data["sheets"] if s["name"] == "Ethernet")
    eth["rows"][2]["cells"]["c1"]["v"] = text  # "Office"
    return data


@pytest.mark.parametrize("headers", [{}, {"X-Role": "viewer"}])
def test_only_admins(client: TestClient, headers) -> None:
    doc_id = _upload(client)
    base = f"/api/documents/{doc_id}/sheet"
    assert client.get(base, headers=headers).status_code == 403
    assert client.post(base, headers=headers).status_code == 403
    assert client.put(base, json={"data": {}, "revision": 1}, headers=headers).status_code == 403
    assert client.get(f"{base}/download", headers=headers).status_code == 403


def test_start_get_save_download(client: TestClient) -> None:
    doc_id = _upload(client)
    assert client.get(f"/api/documents/{doc_id}/sheet", headers=ADMIN).status_code == 404

    sheet = _start(client, doc_id)
    assert (sheet["revision"], sheet["base_version"], sheet["latest_version"]) == (1, 1, 1)
    assert [s["name"] for s in sheet["data"]["sheets"]][:2] == ["Ethernet", "Setup"]

    r = client.put(f"/api/documents/{doc_id}/sheet",
                   json={"data": _edit_first_cell(sheet["data"], "Edited office"), "revision": 1}, headers=ADMIN)
    assert r.status_code == 200, r.text
    assert r.json()["revision"] == 2

    got = client.get(f"/api/documents/{doc_id}/sheet", headers=ADMIN).json()
    assert got["revision"] == 2 and got["updated_by"] == "alice"
    assert got["data"]["sheets"][0]["rows"][2]["cells"]["c1"]["v"] == "Edited office"

    xlsx = client.get(f"/api/documents/{doc_id}/sheet/download?format=xlsx", headers=ADMIN)
    assert xlsx.status_code == 200 and xlsx.headers["content-type"] == XLSX
    ws = load_workbook(io.BytesIO(xlsx.content))["Ethernet"]
    assert (ws["A3"].value, ws["D2"].value) == ("Edited office", "=SUM(D3:D5)")

    pdf = client.get(f"/api/documents/{doc_id}/sheet/download?format=pdf", headers=ADMIN)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert client.app.state.converter.calls[-1] == f"live-{doc_id}.xlsx"


def test_stale_revision_is_rejected(client: TestClient) -> None:
    doc_id = _upload(client)
    sheet = _start(client, doc_id)
    first = client.put(f"/api/documents/{doc_id}/sheet",
                       json={"data": _edit_first_cell(sheet["data"], "mine"), "revision": 1}, headers=ADMIN)
    assert first.status_code == 200
    stale = client.put(f"/api/documents/{doc_id}/sheet",
                       json={"data": _edit_first_cell(sheet["data"], "theirs"), "revision": 1},
                       headers={"X-Role": "network-admin", "X-User": "bob"})
    assert stale.status_code == 409
    assert "alice changed this sheet" in stale.json()["detail"]
    data = client.get(f"/api/documents/{doc_id}/sheet", headers=ADMIN).json()["data"]
    assert data["sheets"][0]["rows"][2]["cells"]["c1"]["v"] == "mine"


def test_invalid_or_huge_sheets_are_rejected(client: TestClient) -> None:
    doc_id = _upload(client)
    _start(client, doc_id)
    bad = {"sheets": [{"id": "s1", "name": "S", "columns": [{"id": "c1"}],
                       "rows": [{"id": "r1", "cells": {"c9": {"v": "x"}}}]}]}
    r = client.put(f"/api/documents/{doc_id}/sheet", json={"data": bad, "revision": 1}, headers=ADMIN)
    assert r.status_code == 422
    huge = b'{"data":' + b" " * (6 * 1024 * 1024) + b"}"
    r = client.put(f"/api/documents/{doc_id}/sheet", content=huge, headers={**ADMIN, "Content-Type": "application/json"})
    assert r.status_code == 413


def test_start_again_and_newer_version(client: TestClient) -> None:
    doc_id = _upload(client)
    sheet = _start(client, doc_id)
    client.put(f"/api/documents/{doc_id}/sheet",
               json={"data": _edit_first_cell(sheet["data"], "edit"), "revision": 1}, headers=ADMIN)
    _upload(client, version=2)

    got = client.get(f"/api/documents/{doc_id}/sheet", headers=ADMIN).json()
    assert (got["base_version"], got["latest_version"]) == (1, 2)  # edits kept, newer version flagged

    again = _start(client, doc_id)
    assert (again["base_version"], again["revision"]) == (2, 3)
    assert again["data"]["sheets"][0]["rows"][2]["cells"]["c1"]["v"] == "Office"


def test_no_ready_version(client: TestClient) -> None:
    client.app.state.converter.fail = True
    r = client.post("/api/documents/upload",
                    files={"file": ("LIS26 _ Event_Network_Requirements_V1.xlsx", sheet_workbook(), XLSX)},
                    headers=ADMIN)
    doc_id = r.json()["document"]["id"]
    assert client.post(f"/api/documents/{doc_id}/sheet", headers=ADMIN).status_code == 409


def test_deleting_the_document_removes_its_sheet(client: TestClient) -> None:
    doc_id = _upload(client)
    _start(client, doc_id)
    client.delete(f"/api/documents/{doc_id}", headers=ADMIN)
    conn = db.connect(client.settings.db_path)
    assert conn.execute("SELECT COUNT(*) FROM live_sheets").fetchone()[0] == 0
    conn.close()
