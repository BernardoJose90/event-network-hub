"""A small workbook with the same patterns as the real Event Network Requirements file."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation


def _fill(hex6: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor="FF" + hex6)


def sheet_workbook() -> bytes:
    wb = Workbook()
    eth = wb.active
    eth.title = "Ethernet"
    setup = wb.create_sheet("Setup")
    wifi = wb.create_sheet("Wi-Fi")
    vlans = wb.create_sheet("Extending Vlans")
    changes = wb.create_sheet("Changes")

    for ws in (eth, setup, wifi, changes):
        ws.sheet_format.defaultColWidth = 14.43

    # Setup: lookup lists. Column D feeds the VLAN dropdowns (also via a named range).
    setup["A1"], setup["D1"] = "Building", "VLANs"
    setup["A1"].font = setup["D1"].font = Font(name="Overpass", size=10, bold=True)
    for i, (building, vlan) in enumerate([("MEO Arena", "General"), ("PAV 1", "Operational_Prod"),
                                          ("PAV 2", "Partner_General")], start=2):
        setup.cell(i, 1, building).font = Font(name="Overpass", size=10)
        setup.cell(i, 4, vlan).font = Font(name="Overpass", size=10)
    wb.defined_names["VLANOptions"] = DefinedName("VLANOptions", attr_text="Setup!$D$2:$D$4")

    # Ethernet: sections with totals, like the real sheet.
    eth.freeze_panes = "A2"
    eth.column_dimensions["A"].width = 32
    rows = [
        ("Purpose", "Location", "VLAN", "Points"),
        ("INTERNAL WORKSPACES", None, None, "=SUM(D3:D4)"),  # too short, like the real H2
        ("Office", "MEO Arena", "General", 2),
        ("Store", "PAV 1", "Operational_Prod", 3),
        ("FoodTrucks", None, "General", 98),  # outside the original range
        ("PARTNER MASTERCLASSES", None, None, "=sum(#REF!)"),
        ("STAGES", None, None, "=SUM(D8)"),
        ("Stage 1", "PAV 2", "General", 5),
        ("Total", None, None, "=sum(D7,D6,D2)"),
    ]
    for r, values in enumerate(rows, start=1):
        for c, value in enumerate(values, start=1):
            cell = eth.cell(r, c, value)
            cell.font = Font(name="Overpass", size=10, bold=r in (1, 2, 6, 7, 9),
                             color="FFFFFFFF" if r in (1, 2, 6, 7, 9) else None)
            if r == 1:
                cell.fill = _fill("222222")
            elif r in (2, 6, 7, 9):
                cell.fill = _fill("666666")
    for formula, cells in (("Setup!A2:A3", "B3:B5"), ("VLANOptions", "C3:C5 C8")):
        dv = DataValidation(type="list", formula1=formula)
        for part in cells.split():
            dv.add(part)
        eth.add_data_validation(dv)

    # Wi-Fi: a merged, filled section row and a fixed-list dropdown.
    wifi.freeze_panes = "A2"
    for c, title in enumerate(["ID", "SSID", "Coverage", "Band", "Purchased"], start=1):
        wifi.cell(1, c, title).font = Font(name="Overpass", size=10, bold=True)
    wifi["A2"] = "Core Networks"
    wifi["A2"].fill = _fill("1D1C1D")
    wifi.merge_cells("A2:D2")
    wifi["A3"], wifi["B3"], wifi["E3"] = 1, "#WebSummit", "Yes"
    for c in range(1, 6):
        wifi.cell(3, c).font = Font(name="Overpass", size=10)
    dv = DataValidation(type="list", formula1='"Yes,No,TBC"')
    dv.add("E3:E5")
    wifi.add_data_validation(dv)

    # Extending Vlans: dropdown + tick-boxes (Google Sheets exports them as booleans).
    vlans["A1"], vlans["B1"], vlans["C1"] = "Vlan name", "FIL", "Altice Arena"
    vlans["A2"], vlans["B2"], vlans["C2"] = "General", True, True
    vlans["A3"], vlans["B3"], vlans["C3"] = None, False, False
    dv = DataValidation(type="list", formula1="Setup!$D$2:$D$4")
    dv.add("A2:A3")
    vlans.add_data_validation(dv)

    # Changes: per-row formulas and one formula the live sheet doesn't support.
    changes["A1"], changes["B1"], changes["G1"], changes["H1"] = "Area", "Change", "V0", "V1"
    changes["A2"], changes["B2"], changes["G2"], changes["H2"] = "Lounge", "=H2-G2", 0, 2
    changes["A3"], changes["B3"], changes["G3"], changes["H3"] = "Screens", "=H3-G3", 21, 0
    changes["C2"] = "=AVERAGE(G2:H2)"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
