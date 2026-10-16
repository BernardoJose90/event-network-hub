from __future__ import annotations

from collections import Counter
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app import db
from app.config import Settings
from app.main import create_app
from app.records.importer import import_workbook, values_missing_from_lists

from .conftest import FakeConverter
from .records_fixture import records_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ADMIN = {"X-Role": "network-admin", "X-User": "alice"}
VIEWER = {"X-Role": "viewer", "X-User": "victor"}


# --- importer -------------------------------------------------------------------


@pytest.fixture(scope="module")
def imported():
    return import_workbook(records_workbook())


def test_venues_groups_and_sections(imported) -> None:
    assert imported.venues == ["Altice Arena", "FIL 1-5"]
    assert imported.groups["wifi"] == ["Core Networks", "Lounges", "Partners"]  # theme-coloured rows included
    assert imported.groups["ethernet"] == ["INTERNAL WORKSPACES", "PARTNER MASTERCLASSES", "Stages"]  # not "Total"


def test_record_counts_and_placement(imported) -> None:
    counts = Counter((r.type, r.venue, r.group) for r in imported.records)
    assert counts[("wifi", "Altice Arena", "Core Networks")] == 2
    assert counts[("wifi", "FIL 1-5", "Partners")] == 1
    assert counts[("ethernet", "Altice Arena", "INTERNAL WORKSPACES")] == 4  # incl. the empty-points Bag Drop row
    assert counts[("ethernet", "Altice Arena", "Stages")] == 2
    assert counts[("partner", None, None)] == 3
    assert counts[("vlan_extension", None, None)] == 2  # the empty row is skipped


def test_values_are_normalised(imported) -> None:
    wifi = next(r for r in imported.records if r.type == "wifi" and r.data.get("ssid") == "#Production")
    assert wifi.data == {"ref_id": "2", "ssid": "#Production", "coverage": "offices", "password_required": "yes",
                         "broadcast": "yes", "vlan": "Operational_Prod", "band": "5 GHz",
                         "password": "StaffPass1!", "purpose": "Staff Network"}
    eth = next(r for r in imported.records if r.type == "ethernet" and r.data.get("purpose") == "Stage 1")
    assert eth.data["ports"] == 13 and eth.data["location"] == "Mystery Hall"
    assert eth.data["bandwidth_ns"] == "50 Mbps" and eth.data["bandwidth_ew"] == "1 Gbps"
    ibm = next(r for r in imported.records if r.type == "partner" and r.data["partner"] == "IBM")
    assert (ibm.data["purchased"], ibm.data["bandwidth_mbps"], ibm.data["pavilion"]) == ("yes", 60, "1")
    vlan = next(r for r in imported.records if r.type == "vlan_extension")
    assert vlan.data == {"vlan": "Operational_Prod", "venues": {"FIL 1-5": True, "Altice Arena": True}}


def test_warnings_and_lists(imported) -> None:
    joined = " | ".join(imported.warnings)
    assert "Scratch notes" in joined
    assert "'maybe' is not a valid broadcast" in joined
    assert "'tbc' is not a number of points" in joined
    assert "'fifty' is not a number" in joined
    assert "placeholder or note row" in joined
    assert imported.lists["vlan"] == ["General", "Operational_Prod", "Operations_Stages", "Partner_Custom"]
    assert imported.lists["bandwidth"] == ["10Mbs", "20Mbs"]
    assert values_missing_from_lists(imported) == {"vlan": ["Operational_Catering"], "location": ["Mystery Hall"]}
    assert len(imported.legacy_log) == 2 and "Executive lounge" in imported.legacy_log[0]
    assert "=D2-C2" not in imported.legacy_log[0]


# --- API ---------------------------------------------------------------------------


@pytest.fixture
def client(tmp_path: Path):
    settings = Settings(data_dir=tmp_path, gotenberg_url="http://unused", gotenberg_timeout_seconds=1,
                        max_upload_bytes=1024 * 1024)
    app = create_app(settings)
    with TestClient(app) as c:
        app.state.converter = FakeConverter()
        c.settings = settings
        yield c


@pytest.fixture
def event(client: TestClient) -> int:
    r = client.post("/api/documents/upload", headers=ADMIN,
                    files={"file": ("LIS26 _ Event_Network_Requirements_V1.xlsx", records_workbook(), XLSX)})
    doc_id = r.json()["document"]["id"]
    r = client.post(f"/api/events/{doc_id}/import", headers=ADMIN, json={})
    assert r.status_code == 200, r.text
    return doc_id


def _records(client, doc, rtype, headers=ADMIN):
    return client.get(f"/api/events/{doc}/records/{rtype}", headers=headers).json()


def _find(client, doc, rtype, **match):
    return next(r for r in _records(client, doc, rtype) if all(r["data"].get(k) == v for k, v in match.items()))


def test_import_result_and_overview(client: TestClient, event: int) -> None:
    ov = client.get(f"/api/events/{event}/overview", headers=VIEWER).json()
    assert ov["has_records"] and ov["counts"] == {"wifi": 4, "ethernet": 7, "partner": 3, "vlan_extension": 2}
    assert ov["flagged_list_values"] == 2
    altice = ov["totals"]["ethernet"][0]
    assert altice["venue"] == "Altice Arena"
    assert [(s["section"], s["ports"]) for s in altice["sections"]] == [
        ("INTERNAL WORKSPACES", 111), ("PARTNER MASTERCLASSES", 0), ("Stages", 13)]
    assert altice["total"] == 124
    ext = _records(client, event, "vlan_extension")[0]["data"]
    venue_ids = {v["name"]: str(v["id"]) for v in ov["venues"]}
    assert ext["venues"] == {venue_ids["FIL 1-5"]: True, venue_ids["Altice Arena"]: True}
    legacy = client.get(f"/api/events/{event}/history?entity=legacy", headers=VIEWER).json()
    assert len(legacy) == 2


def test_import_twice_needs_replace(client: TestClient, event: int) -> None:
    assert client.post(f"/api/events/{event}/import", headers=ADMIN, json={}).status_code == 409
    r = client.post(f"/api/events/{event}/import", headers=ADMIN, json={"replace": True})
    assert r.status_code == 200
    assert client.get(f"/api/events/{event}/overview", headers=ADMIN).json()["counts"]["ethernet"] == 7


def test_passwords_are_masked_and_reveals_logged(client: TestClient, event: int) -> None:
    wifi = _find(client, event, "wifi", ssid="#Production")
    assert "password" not in wifi["data"] and wifi["has_password"] is True
    r = client.post(f"/api/events/{event}/records/wifi/{wifi['id']}/reveal", headers=VIEWER)
    assert r.json() == {"password": "StaffPass1!"}
    entry = client.get(f"/api/events/{event}/history", headers=ADMIN).json()[0]
    assert (entry["action"], entry["user"], entry["record_id"]) == ("reveal", "victor", wifi["id"])


def test_viewers_cannot_change_anything(client: TestClient, event: int) -> None:
    eth = _find(client, event, "ethernet", purpose="Photographers Office")
    base = f"/api/events/{event}"
    attempts = [
        client.post(f"{base}/records/ethernet", headers=VIEWER, json={"data": {}}),
        client.patch(f"{base}/records/ethernet/{eth['id']}", headers=VIEWER, json={"revision": 1, "data": {}}),
        client.delete(f"{base}/records/ethernet/{eth['id']}?revision=1", headers=VIEWER),
        client.post(f"{base}/lists/vlan", headers=VIEWER, json={"value": "X"}),
        client.post(f"{base}/import", headers=VIEWER, json={"replace": True}),
    ]
    assert [a.status_code for a in attempts] == [403] * 5
    assert client.get(f"{base}/records/ethernet", headers={}).status_code == 403  # no role at all


def test_create_update_delete_are_logged(client: TestClient, event: int) -> None:
    ov = client.get(f"/api/events/{event}/overview", headers=ADMIN).json()
    venue = ov["venues"][0]["id"]
    section = next(g["id"] for g in ov["groups"] if g["kind"] == "ethernet" and g["name"] == "Stages")
    r = client.post(f"/api/events/{event}/records/ethernet", headers=ADMIN, json={
        "venue_id": venue, "group_id": section, "data": {"purpose": "Stage 3", "vlan": "General", "ports": 5}})
    assert r.status_code == 201, r.text
    created = r.json()
    totals = client.get(f"/api/events/{event}/overview", headers=ADMIN).json()["totals"]["ethernet"][0]
    assert totals["sections"][2]["ports"] == 18

    r = client.patch(f"/api/events/{event}/records/ethernet/{created['id']}", headers=ADMIN,
                     json={"revision": created["revision"], "data": {"ports": 7, "notes": "BOH"}})
    assert r.status_code == 200 and r.json()["revision"] == created["revision"] + 1
    url = f"/api/events/{event}/records/ethernet/{created['id']}"
    assert client.delete(f"{url}?revision={r.json()['revision']}", headers=ADMIN).status_code == 409  # archive first
    r = client.post(f"{url}/archive", headers=ADMIN, json={"revision": r.json()["revision"]})
    assert r.status_code == 200 and r.json()["archived_by"] == "alice"
    r = client.delete(f"{url}?revision={r.json()['revision']}", headers=ADMIN)
    assert r.status_code == 204

    log = client.get(f"/api/events/{event}/history?entity=ethernet", headers=ADMIN).json()
    assert [(e["action"], e["field"], e["old"], e["new"]) for e in log] == [
        ("delete", None, '{"purpose": "Stage 3", "vlan": "General", "ports": 7, "notes": "BOH"}', None),
        ("archive", None, None, None),
        ("update", "ports", "5", "7"),
        ("update", "notes", None, "BOH"),
        ("create", None, None, None),
    ]
    assert log[2]["summary"] == "Changed Ethernet Points of Ethernet requirement 'Stage 3'"


def test_password_changes_are_logged_without_values(client: TestClient, event: int) -> None:
    wifi = _find(client, event, "wifi", ssid="#Production")
    client.patch(f"/api/events/{event}/records/wifi/{wifi['id']}", headers=ADMIN,
                 json={"revision": wifi["revision"], "data": {"password": "NewSecret9"}})
    entry = client.get(f"/api/events/{event}/history?entity=wifi", headers=ADMIN).json()[0]
    assert (entry["field"], entry["old"], entry["new"]) == ("password", None, None)
    assert "NewSecret9" not in str(client.get(f"/api/events/{event}/history", headers=ADMIN).json())


def test_validation(client: TestClient, event: int) -> None:
    ov = client.get(f"/api/events/{event}/overview", headers=ADMIN).json()
    venue = ov["venues"][0]["id"]
    section = next(g["id"] for g in ov["groups"] if g["kind"] == "ethernet")
    base = f"/api/events/{event}/records/ethernet"

    def create(data, **extra):
        return client.post(base, headers=ADMIN, json={"venue_id": venue, "group_id": section, "data": data, **extra})

    r = create({"purpose": "X", "vlan": "Not_A_Vlan"})
    assert r.status_code == 422 and "not in the vlan list" in r.json()["detail"]
    r = create({"purpose": "X", "ports": "many"})
    assert r.status_code == 422 and r.json()["detail"].startswith("Ethernet Points")
    r = create({"purpose": "X", "unknown_field": 1})
    assert r.status_code == 422
    r = client.post(base, headers=ADMIN, json={"data": {"purpose": "X"}})  # no venue/section
    assert r.status_code == 422
    r = client.post(f"/api/events/{event}/records/partner", headers=ADMIN,
                    json={"venue_id": venue, "data": {"partner": "X"}})
    assert r.status_code == 422


def test_stale_revision_is_rejected(client: TestClient, event: int) -> None:
    eth = _find(client, event, "ethernet", purpose="Photographers Office")
    url = f"/api/events/{event}/records/ethernet/{eth['id']}"
    assert client.patch(url, headers=ADMIN, json={"revision": eth["revision"], "data": {"ports": 3}}).status_code == 200
    r = client.patch(url, headers={"X-Role": "network-admin", "X-User": "bob"},
                     json={"revision": eth["revision"], "data": {"ports": 9}})
    assert r.status_code == 409 and "alice changed this record" in r.json()["detail"]
    assert client.delete(f"{url}?revision={eth['revision']}", headers=ADMIN).status_code == 409


def test_insert_before_and_move(client: TestClient, event: int) -> None:
    rows = [r for r in _records(client, event, "ethernet") if r["data"].get("purpose") in ("Photographers Office", "PCT Event Control")]
    first, second = rows
    r = client.post(f"/api/events/{event}/records/ethernet", headers=ADMIN, json={
        "venue_id": first["venue_id"], "group_id": first["group_id"], "before_id": second["id"],
        "data": {"purpose": "Inserted"}})
    names = [x["data"].get("purpose") for x in _records(client, event, "ethernet") if x["group_id"] == first["group_id"]
             and x["venue_id"] == first["venue_id"]]
    assert names[:3] == ["Photographers Office", "Inserted", "PCT Event Control"]
    inserted = r.json()
    r = client.post(f"/api/events/{event}/records/ethernet/{inserted['id']}/move", headers=ADMIN,
                    json={"revision": inserted["revision"], "before_id": first["id"]})
    assert r.status_code == 200
    names = [x["data"].get("purpose") for x in _records(client, event, "ethernet") if x["group_id"] == first["group_id"]
             and x["venue_id"] == first["venue_id"]]
    assert names[:2] == ["Inserted", "Photographers Office"]


def test_lists(client: TestClient, event: int) -> None:
    base = f"/api/events/{event}/lists"
    lists = client.get(f"/api/events/{event}/overview", headers=ADMIN).json()["lists"]
    general = next(i for i in lists if i["kind"] == "vlan" and i["value"] == "General")
    flagged = next(i for i in lists if i["flagged"])

    assert client.delete(f"{base}/item/{general['id']}", headers=ADMIN).status_code == 409  # in use
    r = client.patch(f"{base}/item/{general['id']}", headers=ADMIN, json={"value": "General_Access"})
    assert r.status_code == 200
    assert _find(client, event, "wifi", ssid="#WebSummit")["data"]["vlan"] == "General_Access"
    assert client.post(f"{base}/vlan", headers=ADMIN, json={"value": "General_Access"}).status_code == 409
    assert client.post(f"{base}/vlan", headers=ADMIN, json={"value": "Guest"}).status_code == 201
    r = client.post(f"{base}/item/{flagged['id']}/accept", headers=ADMIN)
    assert not next(i for i in r.json() if i["id"] == flagged["id"])["flagged"]


def test_venues_and_sections(client: TestClient, event: int) -> None:
    ov = client.get(f"/api/events/{event}/overview", headers=ADMIN).json()
    altice = ov["venues"][0]
    assert client.delete(f"/api/events/{event}/venues/{altice['id']}", headers=ADMIN).status_code == 409
    r = client.post(f"/api/events/{event}/groups/ethernet", headers=ADMIN, json={"name": "Media Areas"})
    media = next(g for g in r.json() if g["name"] == "Media Areas")
    r = client.patch(f"/api/events/{event}/groups/{media['id']}", headers=ADMIN, json={"name": "Media"})
    assert any(g["name"] == "Media" for g in r.json())
    assert client.delete(f"/api/events/{event}/groups/{media['id']}", headers=ADMIN).status_code == 200


def test_deleting_the_document_removes_its_records(client: TestClient, event: int) -> None:
    client.delete(f"/api/documents/{event}", headers=ADMIN)
    conn = db.connect(client.settings.db_path)
    counts = [conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
              for t in ("records", "event_lists", "event_venues", "event_groups", "change_log")]
    conn.close()
    assert counts == [0, 0, 0, 0, 0]


# --- publish ---------------------------------------------------------------------------


def _comparable(client, doc) -> Counter:
    """Records as (type, venue, group, data) with venue/group names, for round-trip checks."""
    ov = client.get(f"/api/events/{doc}/overview", headers=ADMIN).json()
    venues = {v["id"]: v["name"] for v in ov["venues"]}
    groups = {g["id"]: g["name"] for g in ov["groups"]}
    out = Counter()
    for rtype in ("wifi", "ethernet", "partner", "vlan_extension"):
        for r in _records(client, doc, rtype):
            data = dict(r["data"])
            if rtype == "vlan_extension":
                data["venues"] = dict(sorted((venues[int(k)], v) for k, v in data["venues"].items()))
            if rtype == "wifi":
                data["has_password"] = r["has_password"]
            out[(rtype, venues.get(r["venue_id"]), groups.get(r["group_id"]), str(sorted(data.items())))] += 1
    return out


def test_publish_creates_next_version(client: TestClient, event: int) -> None:
    assert client.get(f"/api/events/{event}/overview", headers=ADMIN).json()["unpublished_changes"] == 1  # the import
    r = client.post(f"/api/events/{event}/publish", headers=ADMIN)
    assert r.status_code == 201, r.text
    version = r.json()["version"]
    assert (version["version"], version["status"], version["source"], version["uploaded_by"]) == (2, "ready", "publish", "alice")

    html = client.app.state.converter.html
    assert "LIS26 | Event_Network_Requirements" in html and "Version 2" in html
    assert "#Production" in html and "StaffPass1!" in html  # the published report shows passwords, as before
    assert ">111<" in html and ">124<" in html  # section and venue totals
    assert "Imported requirements from V1" in html

    ov = client.get(f"/api/events/{event}/overview", headers=ADMIN).json()
    assert (ov["unpublished_changes"], ov["latest_version"]) == (0, 2)
    pdf = client.get(f"/api/documents/{event}/download?version=2", headers=VIEWER)
    assert pdf.status_code == 200 and pdf.content == b"%PDF-1.7 fake report"


def test_published_excel_reimports_to_the_same_records(client: TestClient, event: int) -> None:
    before = _comparable(client, event)
    client.post(f"/api/events/{event}/publish", headers=ADMIN)
    xlsx = client.get(f"/api/documents/{event}/download?version=2&format=xlsx", headers=ADMIN)
    assert xlsx.status_code == 200

    r = client.post("/api/documents/upload", headers=ADMIN,
                    files={"file": ("LIS27 _ Event_Network_Requirements_V1.xlsx", xlsx.content, XLSX)})
    other = r.json()["document"]["id"]
    imported = client.post(f"/api/events/{other}/import", headers=ADMIN, json={}).json()
    assert _comparable(client, other) == before
    assert not [w for w in imported["warnings"] if "not recognised" in w]
    lists = lambda d: sorted((i["kind"], i["value"]) for i in client.get(f"/api/events/{d}/overview", headers=ADMIN).json()["lists"])  # noqa: E731
    assert lists(other) == lists(event)


def test_changes_since_publish(client: TestClient, event: int) -> None:
    client.post(f"/api/events/{event}/publish", headers=ADMIN)
    eth = _find(client, event, "ethernet", purpose="Photographers Office")
    client.patch(f"/api/events/{event}/records/ethernet/{eth['id']}", headers=ADMIN,
                 json={"revision": eth["revision"], "data": {"ports": 4}})
    wifi = _find(client, event, "wifi", ssid="#Production")
    client.post(f"/api/events/{event}/records/wifi/{wifi['id']}/reveal", headers=VIEWER)  # not a change
    changes = client.get(f"/api/events/{event}/changes", headers=VIEWER).json()
    assert [c["summary"] for c in changes] == ["Changed Ethernet Points of Ethernet requirement 'Photographers Office'"]
    client.post(f"/api/events/{event}/publish", headers=ADMIN)
    assert "Changed Ethernet Points" in client.app.state.converter.html


def test_publish_rules(client: TestClient, event: int) -> None:
    assert client.post(f"/api/events/{event}/publish", headers=VIEWER).status_code == 403
    r = client.post("/api/documents/upload", headers=ADMIN,
                    files={"file": ("LIS28 _ Event_Network_Requirements_V1.xlsx", records_workbook(), XLSX)})
    empty = r.json()["document"]["id"]
    assert client.post(f"/api/events/{empty}/publish", headers=ADMIN).status_code == 409

    client.app.state.converter.fail = True
    r = client.post(f"/api/events/{event}/publish", headers=ADMIN)
    assert r.status_code == 502 and "Try publishing again" in r.json()["detail"]
    versions = client.get(f"/api/documents/{event}/versions").json()
    assert (versions[0]["version"], versions[0]["status"]) == (2, "failed")
    assert client.get(f"/api/events/{event}/overview", headers=ADMIN).json()["unpublished_changes"] == 1


def test_report_escapes_text(client: TestClient, event: int) -> None:
    eth = _find(client, event, "ethernet", purpose="Photographers Office")
    client.patch(f"/api/events/{event}/records/ethernet/{eth['id']}", headers=ADMIN,
                 json={"revision": eth["revision"], "data": {"notes": "<script>alert(1)</script>"}})
    client.post(f"/api/events/{event}/publish", headers=ADMIN)
    html = client.app.state.converter.html
    assert "<script>alert(1)</script>" not in html and "&lt;script&gt;" in html



def test_archive_hides_and_restore_brings_back(client: TestClient, event: int) -> None:
    eth = _find(client, event, "ethernet", purpose="FoodTrucks / Bar")
    url = f"/api/events/{event}/records/ethernet/{eth['id']}"
    r = client.post(f"{url}/archive", headers=ADMIN, json={"revision": eth["revision"]})
    archived = r.json()

    assert all(x["id"] != eth["id"] for x in _records(client, event, "ethernet"))
    assert [x["id"] for x in client.get(f"/api/events/{event}/records/ethernet?archived=true", headers=VIEWER).json()] == [eth["id"]]
    ov = client.get(f"/api/events/{event}/overview", headers=ADMIN).json()
    assert ov["totals"]["ethernet"][0]["sections"][0]["ports"] == 13  # 111 - 98
    assert (ov["counts"]["ethernet"], ov["archived"]) == (6, {"ethernet": 1})
    client.post(f"/api/events/{event}/publish", headers=ADMIN)
    html = client.app.state.converter.html
    assert "<td>FoodTrucks / Bar</td>" not in html  # not in the tables…
    assert "Archived Ethernet requirement &#x27;FoodTrucks / Bar&#x27;" in html  # …only in the change list

    patch = client.patch(url, headers=ADMIN, json={"revision": archived["revision"], "data": {"ports": 1}})
    assert patch.status_code == 409 and "Restore it" in patch.json()["detail"]
    r = client.post(f"{url}/restore", headers=ADMIN, json={"revision": archived["revision"]})
    assert r.status_code == 200 and r.json()["archived_at"] is None
    assert client.get(f"/api/events/{event}/overview", headers=ADMIN).json()["totals"]["ethernet"][0]["sections"][0]["ports"] == 111
    assert [e["action"] for e in client.get(f"/api/events/{event}/history?entity=ethernet", headers=ADMIN).json()[:2]] == [
        "restore", "archive"]
    assert client.post(f"{url}/archive", headers=VIEWER, json={"revision": 1}).status_code == 403


def test_archived_records_still_block_list_and_section_deletes(client: TestClient, event: int) -> None:
    eth = _find(client, event, "ethernet", purpose="FoodTrucks / Bar")
    client.post(f"/api/events/{event}/records/ethernet/{eth['id']}/archive", headers=ADMIN, json={"revision": eth["revision"]})
    lists = client.get(f"/api/events/{event}/overview", headers=ADMIN).json()["lists"]
    catering = next(i for i in lists if i["value"] == "Operational_Catering")
    assert client.delete(f"/api/events/{event}/lists/item/{catering['id']}", headers=ADMIN).status_code == 409



def test_vlan_extension_ticks_are_logged_by_venue_name(client: TestClient, event: int) -> None:
    ov = client.get(f"/api/events/{event}/overview", headers=ADMIN).json()
    fil = next(str(v["id"]) for v in ov["venues"] if v["name"] == "FIL 1-5")
    ext = _records(client, event, "vlan_extension")[0]
    venues = dict(ext["data"]["venues"]) | {fil: False}
    client.patch(f"/api/events/{event}/records/vlan_extension/{ext['id']}", headers=ADMIN,
                 json={"revision": ext["revision"], "data": {"venues": venues}})
    entry = client.get(f"/api/events/{event}/history?entity=vlan_extension", headers=ADMIN).json()[0]
    assert (entry["summary"], entry["field"], entry["old"], entry["new"]) == (
        "Unticked FIL 1-5 for VLAN extension 'Operational_Prod'", "venue:FIL 1-5", "yes", "no")


def _single_venue_workbook() -> bytes:
    import io

    from openpyxl import Workbook

    from .records_fixture import BLACK, GREY, WHITE_BOLD

    wb = Workbook()
    wifi = wb.active
    wifi.title = "Wi-Fi Requirements"
    wifi.append(["ID", "SSID", "Coverage area", "Password?", "Broadcast?", "VLAN", "Band", "Password value", "Purpose"])
    wifi["A2"] = "Core Networks"
    wifi["A2"].fill, wifi["A2"].font = BLACK, WHITE_BOLD
    wifi.append([1, "#WebSummitVancouver", "all areas", "open network", "yes", "General", "5 GHz", "none", "General"])
    eth = wb.create_sheet("Ethernet Requirements")
    eth.append(["Purpose", "Location", "VLAN", "Area/Room Number", "Type", "Bandwidth per device\nNorth - South",
                "Bandwidth per device\nEast - West", "Capacity", "Ethernet Points", "CCTV", "#Production", "Notes"])
    eth["A2"] = "INTERNAL WORKSPACES"
    eth["A2"].fill, eth["A2"].font = GREY, WHITE_BOLD
    eth.append(["Staff Office", "West Second Floor", "General", "202", "Office", "50 Mbps", "1 Gbps", 10, 4, "no", "yes", "x"])
    settings = wb.create_sheet("0 SETTINGS DO NOT EDIT")
    settings["A1"] = "Named Ranges"
    settings["A3"], settings["B3"], settings["E3"], settings["F3"] = "Building", "Hall/Location/Area", "Bandwidth Allocation", "VLANs"
    settings["B4"], settings["F4"] = "West Second Floor", "General"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_single_venue_workbook_layout() -> None:
    result = import_workbook(_single_venue_workbook())
    assert result.venues == ["Main venue"]
    assert Counter((r.type, r.venue, r.group) for r in result.records) == {
        ("wifi", "Main venue", "Core Networks"): 1, ("ethernet", "Main venue", "INTERNAL WORKSPACES"): 1}
    eth = next(r for r in result.records if r.type == "ethernet")
    assert (eth.data["ports"], eth.data["capacity"], eth.data["notes"]) == (4, "10", "x")
    assert result.lists["location"] == ["West Second Floor"] and result.lists["vlan"] == ["General"]
    note = next(w for w in result.warnings if "columns not imported" in w)
    assert note == "Ethernet Requirements: columns not imported (no matching field): Type, CCTV, #Production"


# --- new events without a workbook --------------------------------------------------


def _new_event(client, code, **body):
    return client.post("/api/events", headers=ADMIN, json={"event_code": code, **body})


def test_create_empty_event_and_publish_v1(client: TestClient) -> None:
    r = _new_event(client, "web26")
    assert r.status_code == 201, r.text
    doc = r.json()
    assert doc["name"] == "WEB26 | Event_Network_Requirements"
    ov = client.get(f"/api/events/{doc['id']}/overview", headers=VIEWER).json()
    assert not ov["has_records"]
    assert [v["name"] for v in ov["venues"]] == ["Main venue"]
    assert [g["name"] for g in ov["groups"] if g["kind"] == "wifi"] == ["Core Networks", "Lounges",
                                                                       "Security Zones", "Partners"]
    assert len([g for g in ov["groups"] if g["kind"] == "ethernet"]) == 10
    assert client.get(f"/api/documents/{doc['id']}/versions").json() == []
    assert client.post(f"/api/events/{doc['id']}/publish", headers=ADMIN).status_code == 409  # nothing yet

    core = next(g for g in ov["groups"] if g["kind"] == "wifi")
    r = client.post(f"/api/events/{doc['id']}/records/wifi", headers=ADMIN,
                    json={"venue_id": ov["venues"][0]["id"], "group_id": core["id"], "data": {"ssid": "WS_Staff"}})
    assert r.status_code == 201, r.text
    r = client.post(f"/api/events/{doc['id']}/publish", headers=ADMIN)
    assert r.status_code == 201, r.text
    versions = client.get(f"/api/documents/{doc['id']}/versions").json()
    assert [(v["version"], v["status"]) for v in versions] == [(1, "ready")]


def test_create_event_rules(client: TestClient, event: int) -> None:
    assert _new_event(client, "LIS26").status_code == 409  # the uploaded fixture's event
    assert _new_event(client, "lis26").status_code == 409
    assert _new_event(client, "L").status_code == 422
    assert _new_event(client, "LIS 26").status_code == 422
    r = client.post("/api/events", headers=VIEWER, json={"event_code": "WEB26"})
    assert r.status_code == 403
    assert _new_event(client, "WEB26", template_id=9999).status_code == 404
    assert _new_event(client, "WEB26", template_id=event).status_code == 404  # an event, not a template
    assert [d["name"] for d in client.get("/api/documents").json()] == ["LIS26 | Event_Network_Requirements"]


def _save_template(client, event, name="LIS", headers=ADMIN):
    return client.post("/api/templates", headers=headers, json={"name": name, "event_id": event})


def test_save_template(client: TestClient, event: int) -> None:
    wifi = _records(client, event, "wifi")
    client.post(f"/api/events/{event}/records/wifi/{wifi[0]['id']}/archive", headers=ADMIN,
                json={"revision": wifi[0]["revision"]})
    r = _save_template(client, event, "lis")
    assert r.status_code == 201, r.text
    tpl = r.json()
    assert (tpl["name"], tpl["kind"]) == ("LIS", "template")

    # only "Start from" lists it; it is not a document
    assert [t["name"] for t in client.get("/api/templates", headers=VIEWER).json()] == ["LIS"]
    assert [d["name"] for d in client.get("/api/documents").json()] == ["LIS26 | Event_Network_Requirements"]

    src = client.get(f"/api/events/{event}/overview", headers=ADMIN).json()
    new = client.get(f"/api/events/{tpl['id']}/overview", headers=ADMIN).json()
    assert new["document"]["kind"] == "template"
    assert [v["name"] for v in new["venues"]] == [v["name"] for v in src["venues"]]
    assert [(g["kind"], g["name"]) for g in new["groups"]] == [(g["kind"], g["name"]) for g in src["groups"]]
    assert [(i["kind"], i["value"]) for i in new["lists"]] == [(i["kind"], i["value"]) for i in src["lists"]]
    assert new["counts"] == {k: v for k, v in src["counts"].items() if k != "partner"}  # active, no partners
    assert all(t["total"] == 0 for t in new["totals"]["ethernet"])
    assert all(r["data"]["ports"] == 0 for r in _records(client, tpl["id"], "ethernet"))
    assert not any(r["has_password"] for r in _records(client, tpl["id"], "wifi"))
    assert _records(client, tpl["id"], "partner") == []
    # other fields are kept
    keep = lambda recs: sorted((r["data"].get("purpose"), r["data"].get("vlan")) for r in recs)
    assert keep(_records(client, tpl["id"], "ethernet")) == keep(_records(client, event, "ethernet"))
    assert client.post(f"/api/events/{tpl['id']}/publish", headers=ADMIN).status_code == 409


def test_template_rules(client: TestClient, event: int) -> None:
    assert _save_template(client, event, headers=VIEWER).status_code == 403
    assert _save_template(client, event, "L").status_code == 422
    assert _save_template(client, 9999).status_code == 404
    tpl = _save_template(client, event).json()
    assert _save_template(client, event, "Lis").status_code == 409
    assert _save_template(client, tpl["id"], "LIS2").status_code == 404  # a template is not an event
    assert client.delete(f"/api/templates/{event}", headers=ADMIN).status_code == 404  # an event
    assert client.delete(f"/api/templates/{tpl['id']}", headers=VIEWER).status_code == 403
    assert client.delete(f"/api/templates/{tpl['id']}", headers=ADMIN).status_code == 204
    assert client.get("/api/templates", headers=ADMIN).json() == []


def test_create_event_from_template(client: TestClient, event: int) -> None:
    tpl = _save_template(client, event).json()
    first = _records(client, tpl["id"], "ethernet")[0]
    client.patch(f"/api/events/{tpl['id']}/records/ethernet/{first['id']}", headers=ADMIN,
                 json={"revision": first["revision"], "data": {"ports": 2}})
    r = _new_event(client, "LIS27", template_id=tpl["id"])
    assert r.status_code == 201, r.text
    doc = r.json()["id"]
    assert _comparable(client, doc) == _comparable(client, tpl["id"])
    assert client.get(f"/api/events/{doc}/overview", headers=ADMIN).json()["totals"]["ethernet"][0]["total"] == 2

    # VLAN extension ticks point at the new event's venues
    venues = {str(v["id"]) for v in client.get(f"/api/events/{doc}/overview", headers=ADMIN).json()["venues"]}
    assert all(set(r["data"]["venues"]) <= venues for r in _records(client, doc, "vlan_extension"))
    assert "template LIS" in client.get(f"/api/events/{doc}/history", headers=ADMIN).json()[0]["summary"]

    # independent: deleting the template keeps the event, changing the event keeps the template
    rec = _records(client, doc, "ethernet")[0]
    client.patch(f"/api/events/{doc}/records/ethernet/{rec['id']}", headers=ADMIN,
                 json={"revision": rec["revision"], "data": {"purpose": "Changed"}})
    assert all(r["data"]["purpose"] != "Changed" for r in _records(client, tpl["id"], "ethernet"))
    client.delete(f"/api/templates/{tpl['id']}", headers=ADMIN)
    assert len(_records(client, doc, "ethernet")) == len(_records(client, event, "ethernet"))


def test_upload_into_created_event(client: TestClient) -> None:
    created = _new_event(client, "WEB26").json()
    r = client.post("/api/documents/upload", headers=ADMIN,
                    files={"file": ("WEB26 _ Event_Network_Requirements_V1.xlsx", records_workbook(), XLSX)})
    assert r.json()["document"]["id"] == created["id"]
