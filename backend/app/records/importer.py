"""Import an Event Network Requirements workbook (.xlsx) into records.

Recognised sheets (names matched case-insensitively):
  "<venue> - Wi-Fi Req", "<venue> Ethernet Req", "Partners Custom Orders",
  "Extending Vlans", "Changes - Log" (kept as a read-only legacy log), "Setup".
Group/section rows are recognised by their dark fill; data rows are the rest.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..xlsx_colors import luminance, resolve, theme_palette
from .models import LIST_FIELDS

WIFI_SHEET = re.compile(r"^(?P<venue>.*?)[\s\-]*wi-?fi\s+req\s*$", re.IGNORECASE)
ETHERNET_SHEET = re.compile(r"^(?P<venue>.*?)[\s\-]*ethernet\s+req\s*$", re.IGNORECASE)
MAX_LEGACY_ROWS = 1000


@dataclass
class ImportedRecord:
    type: str
    data: dict[str, Any]
    venue: str | None = None
    group: str | None = None


@dataclass
class ImportResult:
    venues: list[str] = field(default_factory=list)
    groups: dict[str, list[str]] = field(default_factory=lambda: {"wifi": [], "ethernet": []})
    lists: dict[str, list[str]] = field(default_factory=lambda: {"building": [], "location": [], "vlan": [], "bandwidth": []})
    records: list[ImportedRecord] = field(default_factory=list)
    legacy_log: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def add_unique(self, seq: list[str], value: str) -> None:
        if value not in seq:
            seq.append(value)


def _text(value: Any) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime):
        value = value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ", timespec="minutes")
    elif isinstance(value, date):
        value = value.isoformat()
    text = str(value).strip()
    return text or None


def _number(value: Any) -> float | int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else value
    text = str(value).strip()
    if re.fullmatch(r"-?\d+(\.\d+)?", text):
        n = float(text)
        return int(n) if n.is_integer() else n
    return None


def _is_dark(cell, palette: list[str]) -> bool:
    """Section/group header rows have a dark fill (explicit RGB or a theme colour)."""
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return False
    color = resolve(fill.fgColor, palette)
    return color is not None and luminance(color) < 0.5


def _norm(header: Any) -> str:
    return re.sub(r"\s+", "", str(header or "")).lower()


def _header_map(ws: Worksheet, rules: list[tuple[str, str]], first_column: str | None = None) -> dict[int, str]:
    """Map column index -> field, by header text containing a rule's key (first match wins)."""
    mapping: dict[int, str] = {}
    for cell in ws[1]:
        name = _norm(cell.value)
        for key, target in rules:
            if key in name and target not in mapping.values():
                mapping[cell.column] = target
                break
    if first_column and 1 not in mapping:
        mapping[1] = first_column  # e.g. the Ethernet "Purpose" header that was overwritten with a number
    return mapping


def _row_values(row, mapping: dict[int, str]) -> dict[str, Any]:
    return {mapping[c.column]: c.value for c in row if c.column in mapping}


def _only_first_cell(row) -> bool:
    return all(c.value is None or str(c.value).strip() == "" for c in row[1:])


WIFI_RULES = [("passwordvalue", "password"), ("password?", "password_required"), ("broadcast", "broadcast"),
              ("coverage", "coverage"), ("ssid", "ssid"), ("vlan", "vlan"), ("band", "band"),
              ("purpose", "purpose"), ("id", "ref_id")]
ETHERNET_RULES = [("purpose", "purpose"), ("location", "location"), ("vlan", "vlan"), ("room", "room"),
                  ("north", "bandwidth_ns"), ("east", "bandwidth_ew"), ("capacity", "capacity"),
                  ("ethernetpoints", "ports"), ("points", "ports"), ("notes", "notes")]
PARTNER_RULES = [("partner", "partner"), ("pav", "pavilion"), ("stand", "stand"), ("solution", "solution"),
                 ("bandwidth", "bandwidth_mbps"), ("defaultcable", "default_cables"),
                 ("totalethernetports", "total_ports"), ("wifi", "wifi"), ("vlan", "vlans"),
                 ("note", "notes"), ("purchased", "purchased")]


def _wifi(ws: Worksheet, venue: str, result: ImportResult, palette: list[str]) -> None:
    mapping = _header_map(ws, WIFI_RULES)
    group = None
    skipped = 0
    for row in ws.iter_rows(min_row=2):
        first = row[0]
        if _is_dark(first, palette) and _text(first.value) and _only_first_cell(row):
            group = _text(first.value)
            result.add_unique(result.groups["wifi"], group)
            continue
        raw = _row_values(row, mapping)
        data: dict[str, Any] = {}
        for key, value in raw.items():
            text = _text(value)
            if text is None:
                continue
            if key in ("password_required", "broadcast"):
                low = text.lower()
                allowed = ("yes", "no", "open network") if key == "password_required" else ("yes", "no")
                if low in allowed:
                    data[key] = low
                else:
                    result.warnings.append(f"{ws.title}: '{text}' is not a valid {key.replace('_', ' ')}; left empty")
            else:
                data[key] = text
        if not any(k != "ref_id" for k in data):
            if data or _text(first.value):
                skipped += 1  # placeholder IDs and note rows ("None at this time…")
            continue
        if group is None:
            group = "Networks"
            result.add_unique(result.groups["wifi"], group)
        result.records.append(ImportedRecord("wifi", data, venue, group))
    if skipped:
        result.warnings.append(f"{ws.title}: {skipped} placeholder or note row(s) without network details skipped")


def _ethernet(ws: Worksheet, venue: str, result: ImportResult, palette: list[str]) -> None:
    mapping = _header_map(ws, ETHERNET_RULES, first_column="purpose")
    section = None
    for row in ws.iter_rows(min_row=2):
        first = row[0]
        name = _text(first.value)
        if _is_dark(first, palette) and name:
            if name.lower() == "total":
                continue  # totals are computed
            section = name
            result.add_unique(result.groups["ethernet"], section)
            continue
        raw = _row_values(row, mapping)
        data: dict[str, Any] = {}
        for key, value in raw.items():
            if key == "ports":
                n = _number(value)
                if n is not None:
                    data[key] = n
                elif _text(value):
                    result.warnings.append(f"{ws.title}!{row[0].row}: '{_text(value)}' is not a number of points; left empty")
            elif (text := _text(value)) is not None:
                data[key] = text
        if not data:
            continue
        if section is None:
            section = "Unsorted"
            result.add_unique(result.groups["ethernet"], section)
        result.records.append(ImportedRecord("ethernet", data, venue, section))


def _partners(ws: Worksheet, result: ImportResult) -> None:
    mapping = _header_map(ws, PARTNER_RULES)
    numbers = {"bandwidth_mbps", "default_cables", "total_ports", "wifi", "vlans"}
    for row in ws.iter_rows(min_row=2):
        raw = _row_values(row, mapping)
        data: dict[str, Any] = {}
        for key, value in raw.items():
            if key in numbers:
                n = _number(value)
                if n is not None and n >= 0:
                    data[key] = n
                elif _text(value):
                    result.warnings.append(f"{ws.title}!{row[0].row}: '{_text(value)}' is not a number; left empty")
            elif key == "purchased":
                text = (_text(value) or "").lower()
                if text in ("yes", "no", "tbc"):
                    data[key] = text
            elif (text := _text(value)) is not None:
                data[key] = text
        if data:
            result.records.append(ImportedRecord("partner", data))


def _vlan_extensions(ws: Worksheet, result: ImportResult, venues: list[str]) -> None:
    columns: dict[int, str] = {}
    vlan_col = None
    for cell in ws[1]:
        header = _text(cell.value)
        if not header:
            continue
        if "vlan" in header.lower():
            vlan_col = cell.column
            continue
        venue = next((v for v in venues if v.lower().startswith(header.lower()) or header.lower() in v.lower()), None)
        if venue:
            columns[cell.column] = venue
        else:
            result.warnings.append(f"{ws.title}: column '{header}' doesn't match a venue; not imported")
    for row in ws.iter_rows(min_row=2):
        vlan = _text(row[vlan_col - 1].value) if vlan_col else None
        ticks = {venue: row[col - 1].value is True for col, venue in columns.items()}
        if vlan or any(ticks.values()):
            result.records.append(ImportedRecord("vlan_extension", {"vlan": vlan, "venues": ticks}))


def _setup(ws: Worksheet, result: ImportResult) -> None:
    rules = [("building", "building"), ("hall", "location"), ("location", "location"),
             ("bandwidth", "bandwidth"), ("vlan", "vlan")]
    columns: dict[int, tuple[str, int]] = {}
    for row in ws.iter_rows(min_row=1, max_row=3):
        for cell in row:
            name = _norm(cell.value)
            for key, kind in rules:
                if key in name and cell.column not in columns and kind not in [k for k, _ in columns.values()]:
                    columns[cell.column] = (kind, cell.row)
                    break
    for col, (kind, header_row) in columns.items():
        for (cell,) in ws.iter_rows(min_row=header_row + 1, min_col=col, max_col=col):
            text = _text(cell.value)
            if text and not (cell.font and cell.font.b):
                result.add_unique(result.lists[kind], text)


def _legacy_log(ws: Worksheet, result: ImportResult) -> None:
    headers = [_text(c.value) for c in ws[1]]
    for row in ws.iter_rows(min_row=2, max_row=MAX_LEGACY_ROWS + 1):
        parts = [f"{h}: {_text(c.value)}" for h, c in zip(headers, row) if h and _text(c.value) is not None
                 and not str(c.value).startswith("=")]
        if parts:
            result.legacy_log.append(" · ".join(parts))


def import_workbook(data: bytes) -> ImportResult:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    palette = theme_palette(wb)
    result = ImportResult()
    sheets = [ws for ws in wb.worksheets if isinstance(ws, Worksheet)]

    # Venues first (in sheet order), so the VLAN matrix can refer to them.
    for ws in sheets:
        m = WIFI_SHEET.match(ws.title) or ETHERNET_SHEET.match(ws.title)
        if m and m.group("venue").strip():
            result.add_unique(result.venues, m.group("venue").strip())

    for ws in sheets:
        title = ws.title.strip()
        if (m := WIFI_SHEET.match(title)) and m.group("venue").strip():
            _wifi(ws, m.group("venue").strip(), result, palette)
        elif (m := ETHERNET_SHEET.match(title)) and m.group("venue").strip():
            _ethernet(ws, m.group("venue").strip(), result, palette)
        elif "partner" in title.lower() and "order" in title.lower():
            _partners(ws, result)
        elif "vlan" in title.lower() and "extend" in title.lower():
            _vlan_extensions(ws, result, result.venues)
        elif title.lower() == "setup":
            _setup(ws, result)
        elif "change" in title.lower() and "log" in title.lower():
            _legacy_log(ws, result)
        else:
            result.warnings.append(f"Sheet '{ws.title}' is not recognised; not imported")

    if not result.records:
        result.warnings.append("No requirements were found in this workbook")
    return result


def values_missing_from_lists(result: ImportResult) -> dict[str, list[str]]:
    """Values used by imported records but absent from the imported lists."""
    missing: dict[str, list[str]] = {}
    known = {kind: set(values) for kind, values in result.lists.items()}
    for rec in result.records:
        for field_name, kinds in LIST_FIELDS[rec.type].items():
            value = rec.data.get(field_name)
            if value and not any(value in known.get(k, ()) for k in kinds):
                target = kinds[-1]  # locations go to the location list
                missing.setdefault(target, [])
                if value not in missing[target]:
                    missing[target].append(value)
                known.setdefault(target, set()).add(value)
    return missing
